import datetime as dt
import re

import pytest
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from automacoes import cnj
from automacoes.planilha import gerar_modelo as gm
from automacoes.robos import robo_planilha as rp
from tests.test_modelo_planilha import D, numero, origem_ficticia

HOJE = D("2026-10-06")


class PlanilhaFalsa:
    """Imita a API: valores 'unformatted' (datas como número de série), atualização e inserção de linhas."""

    def __init__(self, xlsx):
        wb = load_workbook(xlsx)
        self.abas = {}
        for ws in wb:
            grade = []
            for row in ws.iter_rows(values_only=True):
                grade.append([rp.serial(v.date() if isinstance(v, dt.datetime) else v)
                              if isinstance(v, (dt.date, dt.datetime)) else ("" if v is None else v) for v in row])
            self.abas[ws.title] = grade
        self.escritas = []

    def ler(self, nome):
        return [list(l) for l in self.abas[nome]]

    def _celula(self, nome, ref, valor):
        letra, linha = re.match(r"([A-Z]+)(\d+)", ref).groups()
        g = self.abas[nome]
        linha, col = int(linha), column_index_from_string(letra)
        while len(g) < linha:
            g.append([])
        while len(g[linha - 1]) < col:
            g[linha - 1].append("")
        g[linha - 1][col - 1] = valor

    @staticmethod
    def _entrada(v):
        """Como o Sheets interpreta USER_ENTERED: apóstrofo = texto literal; AAAA-MM-DD = data."""
        if isinstance(v, str) and v.startswith("'"):
            return v[1:]
        if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            return rp.serial(dt.date.fromisoformat(v))
        return v

    def atualizar(self, nome, dados):
        for d in dados:
            self.escritas.append((nome, d["range"]))
            inicio = d["range"].split(":")[0]
            letra, linha = re.match(r"([A-Z]+)(\d+)", inicio).groups()
            for i, v in enumerate(d["values"][0]):
                self._celula(nome, f"{_letra(column_index_from_string(letra) + i)}{linha}", self._entrada(v))

    def inserir_linhas(self, nome, linha, valores, herdar_formato_de_cima):
        self.escritas.append((nome, f"inserir {linha}"))
        for v in reversed(valores):
            self.abas[nome].insert(linha - 1, [self._entrada(x) for x in v])


def _letra(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


@pytest.fixture
def planilha(tmp_path):
    origem, *_ = origem_ficticia(tmp_path)
    gm.gerar(origem, tmp_path / "v2.xlsx", HOJE)
    return PlanilhaFalsa(tmp_path / "v2.xlsx")


def linhas_prazos(p):
    aba = rp.Aba("Prazos", p.ler("Prazos"), 2)
    return [l for l in aba.linhas if l.get("Processo")], aba


def item(num, disp, texto, orgao="1ª Vara Mista de Mamanguape"):
    return {"numero_processo": cnj.digitos(num), "data_disponibilizacao": disp, "texto": texto, "nomeOrgao": orgao}


def test_insere_publicacao_nova_na_posicao_e_com_formulas(planilha):
    novo = numero("0000009")  # TJPB, Mamanguape (0231)
    robo = rp.Robo(planilha, HOJE)
    robo.processar_djen([item(novo, "2026-10-05", "Intime-se o autor para, no prazo de 15 (quinze) dias, emendar a inicial.")])
    linhas, aba = linhas_prazos(planilha)
    nova = next(l for l in linhas if l["Processo"] == novo)
    assert (nova["Tribunal"], nova["Comarca"], nova["Status"]) == ("TJPB", "Mamanguape", "A triar")
    assert rp.para_data(nova["Publicação"]) == D("2026-10-06")
    assert rp.para_data(nova["Vencimento interno"]) == D("2026-10-27")   # 15 d.u. após 05/10, com 12/10 feriado
    assert rp.para_data(nova["Termo legal"]) == D("2026-10-28")
    assert nova["ID"] == f"{cnj.digitos(novo)}-20261006-manifestação"
    n = nova["_n"]
    assert nova["Criticidade"].startswith(f"=IF(D{n}=")          # fórmulas com o número da própria linha
    tjpb = [l for l in linhas if l["Tribunal"] == "TJPB"]
    assert all(l["_n"] < n or l["Tribunal"] != "TJPB" for l in linhas if l["Tribunal"] == "TJRN")
    assert tjpb[-1]["_n"] >= n                                     # ficou dentro do bloco do TJPB
    proc = rp.Aba("Processos", planilha.ler("Processos"), 1)
    assert any(p.get("Processo") == novo and p.get("Comarca") == "Mamanguape" for p in proc.linhas)
    pubs = rp.Aba("Publicações", planilha.ler("Publicações"), 1)
    assert sum(p.get("Processo") == novo for p in pubs.linhas) == 1
    assert not any(nome == "Prazos" and re.match(rf"{aba.letra['Status']}\d", rng) for nome, rng in planilha.escritas)


def test_nao_duplica(planilha):
    novo = numero("0000009")
    it = item(novo, "2026-10-05", "Manifeste-se no prazo de 5 dias.")
    rp.Robo(planilha, HOJE).processar_djen([it])
    antes = len(linhas_prazos(planilha)[0])
    rp.Robo(planilha, HOJE).processar_djen([it])
    assert len(linhas_prazos(planilha)[0]) == antes
    pubs = rp.Aba("Publicações", planilha.ler("Publicações"), 1)
    assert sum(p.get("Processo") == novo for p in pubs.linhas) == 1


def test_publicacao_ja_migrada_nao_entra_de_novo(planilha):
    n1 = numero("0000001")  # migrado: disponibilizado 01/10, publicado 02/10
    antes = len(linhas_prazos(planilha)[0])
    rp.Robo(planilha, HOJE).processar_djen([item(n1, "2026-10-01", "Decisão. Manifeste-se em 5 dias.")])
    assert len(linhas_prazos(planilha)[0]) == antes


def test_sentenca_sem_prazo_gera_ed_e_apelacao(planilha):
    novo = numero("0000010")
    rp.Robo(planilha, HOJE).processar_djen([item(novo, "2026-10-05", "SENTENÇA. Julgo procedente o pedido. P.R.I.")])
    naturezas = sorted(l["Natureza do prazo"] for l in linhas_prazos(planilha)[0] if l["Processo"] == novo)
    assert naturezas == ["Apelação", "Embargos de declaração"]


def test_completa_linha_manual_e_respeita_ajuste_manual(planilha):
    linhas, aba = linhas_prazos(planilha)
    vazia = max(l["_n"] for l in linhas) + 1
    manual = numero("0000011")
    planilha.atualizar("Prazos", [{"range": f"{aba.letra['Processo']}{vazia}", "values": [[manual]]},
                                  {"range": f"{aba.letra['Disponibilização']}{vazia}", "values": [["2026-10-05"]]},
                                  {"range": f"{aba.letra['Prazo (d.u.)']}{vazia}", "values": [[5]]},
                                  {"range": f"{aba.letra['Status']}{vazia}", "values": [["A triar"]]}])
    robo = rp.Robo(planilha, HOJE)
    robo.completar_e_recalcular()
    l = next(x for x in linhas_prazos(planilha)[0] if x["Processo"] == manual)
    assert rp.para_data(l["Vencimento interno"]) == D("2026-10-13") and l["Comarca"] == "Mamanguape"
    assert robo.rel.completadas == 1 and "calculados pelo robô" in l["Conferência"]


def test_recalcula_quando_feriado_e_conferido(planilha):
    linhas, aba = linhas_prazos(planilha)
    alvo = next(l for l in linhas if l["Status"] == "A triar" and l["Tribunal"] == "TJPB")
    venc = rp.para_data(alvo["Vencimento interno"])
    fer = rp.Aba("Feriados", planilha.ler("Feriados"), 1)
    linha_nova = max(f["_n"] for f in fer.linhas) + 1
    planilha.atualizar("Feriados", [{"range": f"A{linha_nova}:F{linha_nova}",
                                     "values": [[rp.serial(venc), "Ponto facultativo", "TJPB", "FERIADO", "S", "teste"]]}])
    robo = rp.Robo(planilha, HOJE)
    robo.completar_e_recalcular()
    depois = next(l for l in linhas_prazos(planilha)[0] if l["ID"] == alvo["ID"])
    assert rp.para_data(depois["Vencimento interno"]) > venc and "recalculado" in depois["Conferência"]


def test_ensaio_nao_grava(planilha):
    robo = rp.Robo(planilha, HOJE, ensaio=True)
    robo.processar_djen([item(numero("0000012"), "2026-10-05", "Prazo de 5 dias.")])
    robo.registrar_execucao("DJEN")
    assert planilha.escritas == [] and robo.rel.novas == 1 and robo.rel.acoes


def test_registro_de_execucao(planilha):
    robo = rp.Robo(planilha, HOJE)
    robo.registrar_execucao("DJEN")
    valores = planilha.ler("Auditoria")
    marco = next(i for i, l in enumerate(valores) if l and str(l[0]).startswith("Registro de execuções"))
    assert valores[marco + 2][1] == "DJEN" and valores[marco + 2][6] == rp.VERSAO


def test_modelo_alterado_aborta(planilha):
    planilha.abas["Prazos"][1] = [c if c != "Status" else "Situação" for c in planilha.abas["Prazos"][1]]
    with pytest.raises(RuntimeError, match="Status"):
        rp.Robo(planilha, HOJE)


def test_varias_oabs_e_uniao_sem_repetir():
    import configparser
    cfg = configparser.ConfigParser()
    cfg.read_string("[djen]\noab = 29573/PB, 27983/PB; 220690/RJ\nuf = PB\n")
    assert rp.oabs_da_config(cfg["djen"]) == [("29573", "PB"), ("27983", "PB"), ("220690", "RJ")]
    cfg.read_string("[djen]\noab = 29573\nuf = pb\n")
    assert rp.oabs_da_config(cfg["djen"]) == [("29573", "PB")]
    a = {"numero_processo": "1", "data_disponibilizacao": "2026-10-05", "texto": "Intime-se."}
    b = dict(a, texto="Outro ato.")
    assert len(rp.juntar_itens([[a, b], [dict(a)]])) == 2              # mesma comunicação nas duas OABs
    assert len(rp.juntar_itens([[{"id": 7, **a}], [{"id": 7, **b}]])) == 1


def test_listagem_da_leitura(planilha):
    n1 = numero("0000001")
    robo = rp.Robo(planilha, HOJE, ensaio=True)
    robo.processar_djen([item(n1, "2026-10-01", "Decisão."), item(numero("0000013"), "2026-10-05", "Intime-se.")])
    situacoes = {proc: sit for proc, _, _, sit in robo.rel.leitura}
    assert situacoes == {n1: "já na planilha", numero("0000013"): "NOVA"}
