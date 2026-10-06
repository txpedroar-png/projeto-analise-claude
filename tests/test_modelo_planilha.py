import datetime as dt

from openpyxl import Workbook, load_workbook

from automacoes import cnj
from automacoes.planilha import gerar_modelo as gm

D = dt.date.fromisoformat


def numero(seq, trib="15", origem="0231", j="8"):
    return f"{seq}-{cnj.calcular_dv(seq, '2026', j, trib, origem)}.2026.{j}.{trib}.{origem}"


def test_prazos_da_linha():
    assert gm.prazos_da_linha("⚖️ SENTENÇA", "Prazo ED: 18/09/2026 / Apelação: 02/10/2026", "", "Vara") == [
        ("Embargos de declaração", 5, ""), ("Apelação", 15, "prazo recursal presumido (CPC, art. 1.003, § 5º)")]
    assert gm.prazos_da_linha("⚖️ SENTENÇA", "Sentença (ED: 03/09, RI: 11/09)", "", "8º JEC")[1][:2] == ("Recurso inominado", 10)
    assert gm.prazos_da_linha("⏳ PRAZO 10 DIAS", "Manifestação", "", "Vara")[0][:2] == ("Manifestação/providência", 10)
    assert gm.prazos_da_linha("📄 Decisão", "Prazo CPC (5 d.u.): 13/10/2026", "", "Vara")[0][1] == 5
    assert gm.prazos_da_linha("📄 Decisão", "Recurso (15 d.u.): 23/10/2026", "", "Gab.")[0][:2] == ("Recurso", 15)
    assert gm.prazos_da_linha("📄 Publicação", "Ciência de Sentença proferida", "", "Vara") == [("", None, "")]


def test_classificacoes():
    assert gm.instancia("3ª Câmara Cível") == "2º grau" and gm.instancia("8º JEC da Capital") == "Juizado"
    assert gm.area_por_parte("Bradesco Seguros", "0231") == "Bancário – seguros"
    assert gm.area_por_parte("Banco C6", "0231") == "Bancário" and gm.area_por_parte("", "7701") == "Saúde suplementar"
    assert gm.tipo_novo("📄 Publicação", "Ciência de Pauta / Audiência designada") == "Pauta de julgamento"


def origem_ficticia(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Prazos Processuais"
    n1, n2 = numero("0000001"), numero("0000002", "20", "5114")
    linhas = [
        (n1, "Fulana v. Banco X", "1ª Vara Mista de Mamanguape (TJPB)", None, dt.datetime(2026, 10, 2), "📄 Decisão",
         "Prazo CPC (5 d.u.): 09/10/2026", None, None, dt.datetime(2026, 10, 9), None, "Aberto", "Decisão DJe 02/10/2026", None),
        (n1, "Fulana v. Banco X", "1ª Vara Mista de Mamanguape (TJPB)", None, dt.datetime(2026, 10, 2), "📄 Decisão",
         "Prazo CPC (5 d.u.): 09/10/2026", None, None, dt.datetime(2026, 10, 9), None, "Aberto", "Decisão DJe 02/10/2026", None),
        (n2, None, "2ª Vara de Canguaretama (TJRN)", None, dt.datetime(2026, 2, 20), "⏳ PRAZO 5 DIAS",
         "Manifestação / Providência no prazo de 5 dias", None, None, dt.datetime(2026, 2, 27), None, "Intempestivo",
         "Solicitados extratos ao cliente", None),
        (n2, None, "2ª Vara de Canguaretama (TJRN)", None, dt.datetime(2026, 2, 16), "📄 Publicação",
         "Publicação / Intimação simples", None, None, "-", None, "Sem Prazo", None, None),
    ]
    for r, linha in enumerate(linhas, 6):
        for c, v in enumerate(linha, 1):
            ws.cell(r, c, v)
    t = wb.create_sheet("Triagem_Diaria")
    t.append(["Data Publicação", "Unidade", "Processo", "Alerta", "Resumo"])
    t.append([dt.datetime(2026, 10, 2), "1ª Vara Mista de Mamanguape (TJPB)", n1, "📄 Decisão",
              "Publicada Decisão no DJe em 02/10/2026 (Disponibilizado em 01/10/2026)"])
    caminho = tmp_path / "origem.xlsx"
    wb.save(caminho)
    return caminho, n1, n2


def test_migracao_e_modelo(tmp_path):
    origem, n1, n2 = origem_ficticia(tmp_path)
    resumo = gm.gerar(origem, tmp_path / "v2.xlsx", D("2026-10-06"))
    assert resumo["prazos"] == 3 and resumo["processos"] == 2  # duplicata removida
    wb = load_workbook(tmp_path / "v2.xlsx")
    assert wb.sheetnames[:2] == ["Painel", "Prazos"] and "Vista TJPB" in wb.sheetnames
    ws = wb["Prazos"]
    linhas = {(ws.cell(r, 4).value, ws.cell(r, 7).value.date() if ws.cell(r, 7).value else None): r for r in range(3, 6)}
    r = linhas[(n1, D("2026-10-02"))]
    assert ws.cell(r, 1).value == "TJPB" and ws.cell(r, 2).value == "Mamanguape"
    assert ws.cell(r, 6).value.date() == D("2026-10-01")       # disponibilização lida da triagem
    assert ws.cell(r, 10).value.date() == D("2026-10-08")      # interno: 02,05,06,07,08
    assert ws.cell(r, 11).value.date() == D("2026-10-09")      # termo legal
    assert ws.cell(r, 16).value == "A triar"
    r = linhas[(n2, D("2026-02-20"))]
    assert ws.cell(r, 16).value == "A verificar (migração)"
    assert ws.cell(r, 19).value == "Solicitados extratos ao cliente"  # nota humana vai para Observações
    assert "disponibilização deduzida" in ws.cell(r, 23).value
    r = linhas[(n2, D("2026-02-16"))]
    assert ws.cell(r, 16).value == "Sem providência (ciência)"
    assert ws.cell(3, 13).value.startswith("=IF(D3") and "NETWORKDAYS" in ws.cell(3, 12).value
    proc = wb["Processos"]
    clientes = {proc.cell(r, 4).value: proc.cell(r, 5).value for r in range(2, 4)}
    assert clientes[n1] == "Fulana"
