import datetime as dt

from openpyxl import Workbook, load_workbook

from automacoes.planilha import cadastro as cad
from automacoes.planilha import gerar_modelo as gm
from tests.test_modelo_planilha import D, origem_ficticia


def cpf(base9: str) -> str:
    d = base9
    for n in (9, 10):
        soma = sum(int(d[i]) * (n + 1 - i) for i in range(n))
        d += str(soma * 10 % 11 % 10)
    return cad.formatar_cpf(d)


CAB = ["Data Extração", "Nome do Cliente", "CPF", "Prioridade", "Classificação / Comarca", "Espécie da Demanda",
       "Ação Sugerida", "Pendências Probatórias?", "Link da Ficha", "Link da Pasta"]


def cadastro_ficticio(tmp_path):
    wb = Workbook()
    painel = wb.active
    painel.title = "Painel de Controle"
    painel.append(["PAINEL EXECUTIVO"])
    painel.append(["08/09/2026 21:46", "FULANO DE TAL", cpf("111222333"), "Idoso 70 anos",
                   "Ações Bancárias > Paraíba > Pedras de Fogo", "RMC/RCC", "Ação declaratória", "SIM ⚠️",
                   "https://docs.google.com/document/d/abc", "https://drive.google.com/drive/folders/xyz"])
    base = wb.create_sheet("Página1")
    base.append(CAB)
    base.append(["02/09/2026 23:42", "Beltrana da Silva", cpf("444555666"), "Não identificada",
                 "Araújo, Azevedo e Costa > Consumidor > Beltrana vs. LATAM", "Extravio de bagagem", "Indenizatória",
                 "OK ✅", '=HYPERLINK("https://docs.google.com/document/d/f1","📄 Abrir Ficha")',
                 '=HYPERLINK("https://drive.google.com/drive/folders/p1","📁 Abrir Pasta")'])
    base.append(["02/09/2026 23:42", "Ciclano Souza", "123.456.789-00", "", "Ações Bancárias > Paraíba > Mamanguape",
                 "Consignado", "Declaratória", "SIM ⚠️", "", ""])
    base.append(["01/09/2026 10:00", "Beltrana da Silva (antiga)", cpf("444555666"), "", "Consumidor", "", "", "", "", ""])
    caminho = tmp_path / "cadastro.xlsx"
    wb.save(caminho)
    return caminho


def test_cpf():
    assert cad.cpf_valido(cpf("111222333").replace(".", "").replace("-", ""))
    assert not cad.cpf_valido("12345678900") and not cad.cpf_valido("11111111111")


def test_ler_cadastro(tmp_path):
    clientes, avisos = cad.ler_cadastro(cadastro_ficticio(tmp_path))
    por_nome = {c["Nome"]: c for c in clientes}
    assert set(por_nome) == {"FULANO DE TAL", "Beltrana da Silva", "Ciclano Souza"}  # duplicata antiga descartada
    fulano = por_nome["FULANO DE TAL"]
    assert "fora da base" in fulano["Conferência"] and fulano["Comarca provável"] == "Pedras de Fogo"
    assert fulano["Área"] == "Bancário" and fulano["Situação documental"] == "Pendente de documentos"
    beltrana = por_nome["Beltrana da Silva"]
    assert beltrana["Área"] == "Aeronáutico" and beltrana["Situação documental"] == "Pronto para ajuizamento"
    assert beltrana["Ficha"] == "https://docs.google.com/document/d/f1"
    assert "inválido" in por_nome["Ciclano Souza"]["Conferência"]
    assert any("repetido" in a for a in avisos)


def test_modelo_com_clientes_e_vinculo(tmp_path):
    origem, n1, _ = origem_ficticia(tmp_path)   # processo n1 tem cliente "Fulana"
    caminho = cadastro_ficticio(tmp_path)
    wb = load_workbook(caminho)
    wb["Página1"].append(["03/09/2026 09:00", "Fulana", cpf("777888999"), "", "Ações Bancárias > Paraíba > Mamanguape",
                          "", "", "OK ✅", "", ""])
    wb.save(caminho)
    resumo = gm.gerar(origem, tmp_path / "v2.xlsx", D("2026-10-06"), caminho)
    assert resumo["clientes"] == 4
    v2 = load_workbook(tmp_path / "v2.xlsx")
    assert "Clientes" in v2.sheetnames
    proc = v2["Processos"]
    sugestao = {proc[f"D{r}"].value: proc[f"{gm.PC['Sugestão de vínculo (conferir)']}{r}"].value for r in (2, 3)}
    assert sugestao[n1] is None               # nome de uma palavra só nunca gera sugestão
    assert proc[f"{gm.PC['Vínculo com o cadastro']}2"].value.startswith("=IF(D2")
    cli = v2["Clientes"]
    assert cli[f"{gm.CC['Etapa']}2"].value.startswith("=IF(A2")


def test_sugestao_por_nome_completo():
    clientes = [{"Nome": "JOSÉ NILTON DE OLIVEIRA", "CPF": cpf("123123123")},
                {"Nome": "Maria da Silva", "CPF": cpf("321321321")}, {"Nome": "MARIA DA SILVA", "CPF": cpf("456456456")}]
    assert cad.sugerir_vinculo("Jose Nilton de Oliveira", clientes)["CPF"] == cpf("123123123")
    assert cad.sugerir_vinculo("Maria da Silva", clientes) is None    # homônimos: não sugere
    assert cad.sugerir_vinculo("Nilton", clientes) is None
