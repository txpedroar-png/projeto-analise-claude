"""Atualização de descontos indevidos: correção por série do Bacen (SGS) e juros de mora.

Correções em relação à versão anterior:
- mês sem índice publicado gerava correção 0% em silêncio; agora a linha é recusada;
- data de desconto posterior à data-base travava o programa (laço infinito);
- "1.500" (mil e quinhentos) era lido como 1,5;
- a data-base era "hoje" (o resultado mudava a cada execução); agora é fixa e registrada;
- repetição em dobro (art. 42, parágrafo único, do CDC) disponível por opção;
- memorial: abas Cálculo, Índices (série mês a mês) e Premissas (todos os parâmetros).

Convenções (as mesmas da versão anterior, agora declaradas na aba Premissas):
- correção: produto de (1 + índice) do mês do desconto ao mês anterior à data-base;
- juros simples sobre o valor corrigido, por meses civis inteiros entre o termo inicial e a
  data-base; termo inicial = citação (art. 405 do CC) ou desconto (Súmula 54 do STJ).
O título judicial prevalece sobre qualquer padrão daqui.

Uso:
    python -m automacoes.robos.calculadora_bacen Calculo_Valores.xlsx --citacao 28/05/2019 --dobra
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook

SERIES = {"INPC": 188, "IPCA": 433}
SGS_URL = "https://api.bcb.gov.br/dados/serie/bcdata.sgs.{codigo}/dados"
VIGENCIA_LEI_14905 = dt.date(2024, 8, 30)
Mes = tuple[int, int]


@dataclass
class Parametros:
    indice: str
    data_base: dt.date
    citacao: dt.date | None
    termo_juros: str       # "citacao" | "desconto"
    taxa_juros_mensal: float  # 0 = sem juros
    dobra: bool


def ler_data(v) -> dt.date:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data ilegível: {v!r}")


def ler_valor(v) -> float:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).upper().replace("R$", "").replace(" ", "").replace(" ", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")  # 1.500 = mil e quinhentos
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"valor ilegível: {v!r}") from None


def mes(d: dt.date) -> Mes:
    return (d.year, d.month)


def proximo(m: Mes) -> Mes:
    return (m[0] + 1, 1) if m[1] == 12 else (m[0], m[1] + 1)


def meses_civis(inicio: dt.date, fim: dt.date) -> int:
    return max(0, (fim.year - inicio.year) * 12 + fim.month - inicio.month)


def baixar_serie(codigo: int, desde: dt.date, ate: dt.date) -> dict[Mes, float]:
    params = {"formato": "json", "dataInicial": f"{desde:%d/%m/%Y}", "dataFinal": f"{ate:%d/%m/%Y}"}
    resp = requests.get(SGS_URL.format(codigo=codigo), params=params, timeout=30)
    if resp.status_code != 200:
        raise RuntimeError(f"Bacen SGS {codigo}: HTTP {resp.status_code} {resp.text[:200]}")
    return {mes(ler_data(i["data"])): float(i["valor"]) / 100 for i in resp.json()}


def data_base_padrao(indices: dict[Mes, float]) -> dt.date:
    """1º dia do mês seguinte ao último índice publicado: todos os meses têm índice."""
    a, m = proximo(max(indices))
    return dt.date(a, m, 1)


def fator_correcao(desconto: dt.date, data_base: dt.date, indices: dict[Mes, float]) -> float:
    if desconto > data_base:
        raise ValueError(f"desconto ({desconto:%d/%m/%Y}) posterior à data-base")
    fator, m, fim = 1.0, mes(desconto), mes(data_base)
    faltando = []
    while m < fim:
        if m not in indices:
            faltando.append(f"{m[1]:02d}/{m[0]}")
        else:
            fator *= 1 + indices[m]
        m = proximo(m)
    if faltando:
        raise ValueError(f"sem índice publicado para {', '.join(faltando)}")
    return fator


def colunas(p: Parametros) -> list[str]:
    return ["Valor base" + (" (dobro)" if p.dobra else ""), "Fator de correção", f"Valor corrigido ({p.indice})",
            "Meses de juros", "Juros", "Total", "Observação"]


def calcular_linha(data, valor, p: Parametros, indices) -> dict:
    desconto, original = ler_data(data), ler_valor(valor)
    base = original * (2 if p.dobra else 1)
    fator = fator_correcao(desconto, p.data_base, indices)
    corrigido = round(base * fator, 2)
    termo = desconto if p.termo_juros == "desconto" else max(p.citacao, desconto)
    meses = meses_civis(termo, p.data_base) if p.taxa_juros_mensal else 0
    juros = round(corrigido * meses * p.taxa_juros_mensal, 2)
    return dict(zip(colunas(p), [round(base, 2), round(fator, 8), corrigido, meses, juros,
                                 round(corrigido + juros, 2), ""]))


def premissas(p: Parametros, codigo: int, n_ok: int, n_erro: int) -> list[tuple[str, str]]:
    linhas = [
        ("Data-base", f"{p.data_base:%d/%m/%Y}"),
        ("Índice de correção", f"{p.indice} (Bacen SGS {codigo}), consultado em {dt.date.today():%d/%m/%Y}"),
        ("Convenção da correção", "índice do mês do desconto até o mês anterior à data-base"),
        ("Repetição em dobro", "Sim: art. 42, parágrafo único, do CDC" if p.dobra else "Não"),
        ("Juros de mora", f"{p.taxa_juros_mensal:.4%} ao mês, simples, meses civis inteiros"
         if p.taxa_juros_mensal else "Não aplicados"),
        ("Termo inicial dos juros", "data de cada desconto (Súmula 54 do STJ)" if p.termo_juros == "desconto"
         else f"citação em {p.citacao:%d/%m/%Y} ou desconto, o que for posterior (art. 405 do CC)"),
        ("Linhas calculadas / recusadas", f"{n_ok} / {n_erro}"),
        ("Revisão", "Minuta: conferir com o título judicial antes de juntar aos autos."),
    ]
    if p.taxa_juros_mensal and p.data_base > VIGENCIA_LEI_14905:
        linhas.append(("ATENÇÃO", "Data-base após 30/08/2024 (Lei 14.905/2024): sem previsão no título, "
                                  "a taxa legal do art. 406 do CC substitui 1% a.m. Conferir."))
    return linhas


def ler_planilha(entrada: Path) -> tuple[list[str], list[tuple], dict[str, int]]:
    linhas = list(load_workbook(entrada, data_only=True).active.iter_rows(values_only=True))
    cab = [str(c or "").strip() for c in linhas[0]]
    idx = {c.lower(): i for i, c in enumerate(cab)}
    if "data" not in idx or "valor" not in idx:
        raise ValueError(f"A primeira linha precisa ter as colunas 'Data' e 'Valor' (achei: {cab})")
    dados = [r for r in linhas[1:] if r[idx["data"]] not in (None, "") or r[idx["valor"]] not in (None, "")]
    return cab, dados, idx


def desconto_mais_antigo(dados, idx) -> dt.date:
    datas = []
    for r in dados:
        try:
            datas.append(ler_data(r[idx["data"]]))
        except ValueError:
            pass
    return min(datas, default=dt.date.today())


def processar(cab, dados, idx, saida: Path, p: Parametros, indices: dict[Mes, float], codigo: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cálculo"
    novas = colunas(p)
    ws.append(cab + novas)
    n_ok = n_erro = 0
    total = 0.0
    for row in dados:
        try:
            res = calcular_linha(row[idx["data"]], row[idx["valor"]], p, indices)
            n_ok += 1
            total += res["Total"]
        except (ValueError, TypeError) as e:
            res = {"Observação": f"RECUSADA: {e}"}
            n_erro += 1
        ws.append(list(row) + [res.get(c, "") for c in novas])
    rodape = [""] * (len(cab) + len(novas))
    rodape[0], rodape[len(cab) + novas.index("Total")] = "TOTAL", round(total, 2)
    ws.append([])
    ws.append(rodape)

    wi = wb.create_sheet("Índices")
    wi.append(["Mês", f"{p.indice} (%)"])
    for m in sorted(indices):
        if m < mes(p.data_base):
            wi.append([f"{m[1]:02d}/{m[0]}", round(indices[m] * 100, 4)])
    wp = wb.create_sheet("Premissas")
    for linha in premissas(p, codigo, n_ok, n_erro):
        wp.append(linha)
    wb.save(saida)
    return n_ok, n_erro, round(total, 2)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Atualiza descontos indevidos com índice do Bacen e juros de mora.")
    p.add_argument("entrada", type=Path, nargs="?", default=Path("Calculo_Valores.xlsx"))
    p.add_argument("-o", "--saida", type=Path, default=Path("Planilha_Atualizada_Pronta.xlsx"))
    p.add_argument("--indice", choices=sorted(SERIES), default="INPC")
    p.add_argument("--data-base", type=ler_data, help="padrão: mês seguinte ao último índice publicado")
    p.add_argument("--citacao", type=ler_data, help="DD/MM/AAAA (obrigatória se termo de juros = citação)")
    p.add_argument("--termo-juros", choices=["citacao", "desconto"], default="citacao")
    p.add_argument("--juros", type=float, default=1.0, help="%% ao mês; 0 = sem juros (padrão 1)")
    p.add_argument("--dobra", action="store_true", help="repetição em dobro (art. 42, parágrafo único, do CDC)")
    a = p.parse_args(argv)
    if a.juros and a.termo_juros == "citacao" and not a.citacao:
        p.error("informe --citacao ou use --termo-juros desconto")

    codigo = SERIES[a.indice]
    cab, dados, idx = ler_planilha(a.entrada)
    inicio = desconto_mais_antigo(dados, idx).replace(day=1)
    indices = baixar_serie(codigo, inicio, a.data_base or dt.date.today())
    params = Parametros(a.indice, a.data_base or data_base_padrao(indices), a.citacao, a.termo_juros, a.juros / 100, a.dobra)
    n_ok, n_erro, total = processar(cab, dados, idx, a.saida, params, indices, codigo)
    brl = f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    print(f"Data-base {params.data_base:%d/%m/%Y} | {n_ok} linhas calculadas | {n_erro} recusadas | "
          f"total R$ {brl} | arquivo: {a.saida}")
    return 1 if n_erro else 0


if __name__ == "__main__":
    raise SystemExit(main())
