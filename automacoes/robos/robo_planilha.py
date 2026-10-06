"""Robô de prazos para a Planilha Google (modelo v2).

A cada execução:
1. lê as abas Prazos, Processos, Feriados e Comarcas da planilha;
2. completa linhas digitadas pela equipe (disponibilização + prazo, sem vencimento) e recalcula os
   prazos em aberto quando o calendário mudou (ex.: feriado marcado como conferido);
3. consulta o DJEN pela OAB e insere cada publicação nova uma única vez, na posição
   tribunal -> comarca -> vencimento, com as mesmas fórmulas do modelo;
4. cadastra processos novos em Processos e registra a publicação em Publicações;
5. registra a execução na aba Auditoria.

Nunca escreve em Status, Data do protocolo, Observações, Link, CPF nem em linhas encerradas.
Linha com "ajuste manual" nas Observações não é recalculada.

Uso:
    python -m automacoes.robos.robo_planilha --ensaio     # mostra o que faria, sem gravar
    python -m automacoes.robos.robo_planilha
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import time
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from automacoes import cnj
from automacoes.planilha.gerar_modelo import (
    ORDEM_TRIBUNAIS, STATUS_ENCERRADOS, formulas_processos, formulas_prazos, identificador, instancia,
)
from automacoes.prazos import Calendario, calcular_prazo
from automacoes.robos.comum import carregar_config, configurar_log, log
from automacoes.triagem import detectar_prazo, tipo_ato

VERSAO = "robo_planilha 1.0"
ATIVOS = ("A triar", "Em elaboração", "Aguardando cliente/documento", "Minuta em revisão")
OBRIGATORIAS = ["Tribunal", "Comarca", "Órgão julgador", "Processo", "Disponibilização", "Publicação",
                "Prazo (d.u.)", "Natureza do prazo", "Vencimento interno", "Termo legal", "Tipo de ato",
                "Providência", "Status", "Conferência", "Teor / resumo do ato", "Fonte", "ID",
                "Cliente", "Área / nicho", "Dias úteis restantes", "Criticidade", "Data do protocolo",
                "Folga no protocolo (d.u.)", "Observações da equipe"]
EPOCA = dt.date(1899, 12, 30)
LINHA_CAB_PRAZOS, LINHA_CAB_OUTRAS = 2, 1


# ------------------------------------------------------------------ conversões
def serial(d: dt.date) -> int:
    return (d - EPOCA).days


def para_data(v) -> dt.date | None:
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return EPOCA + dt.timedelta(days=int(v))
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def para_int(v) -> int | None:
    try:
        n = int(float(str(v).replace(",", ".")))
        return n if n > 0 else None
    except (TypeError, ValueError):
        return None


def celula(v):
    """Valor pronto para USER_ENTERED. Datas em AAAA-MM-DD (lidas como data em qualquer idioma da planilha);
    texto que pareça número, data ou fórmula vai como texto literal."""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str) and (v[:1] in "=+-@" or re.fullmatch(r"[\d.,/\s-]+", v or "x")):
        return "'" + v
    return "" if v is None else v


# ------------------------------------------------------------------ planilha
class Aba:
    """Cópia local de uma aba: cabeçalho -> letra, linhas como dicionários com o número da linha."""

    def __init__(self, nome: str, valores: list[list], linha_cabecalho: int):
        self.nome = nome
        self.linha_cab = linha_cabecalho
        cab = valores[linha_cabecalho - 1] if len(valores) >= linha_cabecalho else []
        self.letra = {str(c).strip(): get_column_letter(i) for i, c in enumerate(cab, 1) if str(c).strip()}
        self.ncols = len(cab)
        self.linhas = []
        for n, linha in enumerate(valores[linha_cabecalho:], start=linha_cabecalho + 1):
            self.linhas.append({"_n": n, **{str(c).strip(): (linha[i] if i < len(linha) else "")
                                           for i, c in enumerate(cab) if str(c).strip()}})

    def faltando(self, nomes) -> list[str]:
        return [n for n in nomes if n not in self.letra]

    def deslocar(self, a_partir: int) -> None:
        for linha in self.linhas:
            if linha["_n"] >= a_partir:
                linha["_n"] += 1


@dataclass
class Relatorio:
    novas: int = 0
    completadas: int = 0
    recalculadas: int = 0
    processos_novos: int = 0
    publicacoes: int = 0
    para_revisao: int = 0
    lidas: int = 0
    erros: list[str] = field(default_factory=list)
    acoes: list[str] = field(default_factory=list)


class Robo:
    def __init__(self, planilha, hoje: dt.date, ensaio: bool = False):
        self.p, self.hoje, self.ensaio = planilha, hoje, ensaio
        self.rel = Relatorio()
        self.prazos = Aba("Prazos", planilha.ler("Prazos"), LINHA_CAB_PRAZOS)
        if falta := self.prazos.faltando(OBRIGATORIAS):
            raise RuntimeError(f"Aba Prazos sem as colunas {falta}: o modelo foi alterado; nada foi gravado.")
        self.processos = Aba("Processos", planilha.ler("Processos"), LINHA_CAB_OUTRAS)
        self.publicacoes = Aba("Publicações", planilha.ler("Publicações"), LINHA_CAB_OUTRAS)
        feriados = Aba("Feriados", planilha.ler("Feriados"), LINHA_CAB_OUTRAS)
        self.registros_feriados = [
            {"data": para_data(f["Data"]), "descricao": f.get("Descrição", ""), "abrangencia": f.get("Abrangência", ""),
             "tipo": f.get("Tipo", ""), "conferido": f.get("Conferido", "")}
            for f in feriados.linhas if para_data(f.get("Data"))]
        comarcas = Aba("Comarcas", planilha.ler("Comarcas"), LINHA_CAB_OUTRAS)
        self.comarcas = {(str(c["Tribunal"]), str(c["Código de origem"]).zfill(4)): str(c["Comarca"])
                         for c in comarcas.linhas if c.get("Tribunal")}
        self.calendarios: dict = {}

    # ---------------------------------------------------------- utilidades
    def calendario(self, tribunal: str, comarca_cod: str) -> Calendario:
        chave = (tribunal, comarca_cod)
        if chave not in self.calendarios:
            self.calendarios[chave] = Calendario.de_registros(self.registros_feriados, tribunal or "TJPB", comarca_cod)
        return self.calendarios[chave]

    def comarca(self, tribunal: str, cod: str) -> str:
        return self.comarcas.get((tribunal, cod), f"Cód. {cod}")

    def gravar(self, aba: Aba, linha: int, valores: dict) -> None:
        """Atualiza só as colunas indicadas da linha."""
        dados = [{"range": f"{aba.letra[nome]}{linha}", "values": [[celula(v)]]} for nome, v in valores.items()]
        self.rel.acoes.append(f"{aba.nome}!{linha}: " + ", ".join(f"{k}={v}" for k, v in valores.items()))
        if not self.ensaio:
            self.p.atualizar(aba.nome, dados)

    def linha_completa(self, aba: Aba, valores: dict, formulas: dict) -> list:
        linha = [""] * aba.ncols
        for nome, v in valores.items():
            if nome in aba.letra:
                linha[_indice(aba.letra[nome])] = celula(v)
        for letra, formula in formulas.items():
            linha[_indice(letra)] = formula
        return linha

    # ---------------------------------------------------------- 1. completar e recalcular
    def completar_e_recalcular(self) -> None:
        for linha in self.prazos.linhas:
            numero, status = str(linha.get("Processo") or ""), str(linha.get("Status") or "")
            disp, dias = para_data(linha.get("Disponibilização")), para_int(linha.get("Prazo (d.u.)"))
            if not numero or not disp or not dias or status in STATUS_ENCERRADOS:
                continue
            if "ajuste manual" in str(linha.get("Observações da equipe") or "").lower():
                continue
            venc_atual = para_data(linha.get("Vencimento interno"))
            if venc_atual and status not in ATIVOS:
                continue
            d = cnj.digitos(numero)
            trib = cnj.tribunal(d) or str(linha.get("Tribunal") or "")
            r = calcular_prazo(disp, dias, self.calendario(trib, cnj.comarca(d)), self.hoje)
            interno, legal = dt.date.fromisoformat(r.vencimento_interno), dt.date.fromisoformat(r.termo_legal)
            if venc_atual == interno and para_data(linha.get("Termo legal")) == legal:
                continue
            valores = {"Vencimento interno": interno, "Termo legal": legal}
            if not venc_atual:
                self.rel.completadas += 1
                if not linha.get("Tribunal"):
                    valores["Tribunal"] = trib
                if not linha.get("Comarca"):
                    valores["Comarca"] = self.comarca(trib, cnj.comarca(d))
                if not linha.get("Publicação"):
                    valores["Publicação"] = self.calendario(trib, cnj.comarca(d)).somar_dias_uteis(disp, 1)
                if not linha.get("Fonte"):
                    valores["Fonte"] = "Manual (completado pelo robô)"
                if not linha.get("ID"):
                    pub = valores.get("Publicação") or para_data(linha.get("Publicação"))
                    valores["ID"] = identificador(d, pub, str(linha.get("Natureza do prazo") or ""))
                nota = f"vencimentos calculados pelo robô em {self.hoje:%d/%m/%Y}"
            else:
                self.rel.recalculadas += 1
                nota = (f"recalculado em {self.hoje:%d/%m/%Y} (calendário alterado): interno era "
                        f"{venc_atual:%d/%m/%Y}")
            valores["Conferência"] = "; ".join(filter(None, [str(linha.get("Conferência") or ""), nota]))
            self.gravar(self.prazos, linha["_n"], valores)
            linha.update(valores)

    # ---------------------------------------------------------- 2. publicações novas
    def prazos_do_ato(self, texto: str, tipo: str, orgao: str) -> list[tuple[str, int | None, list[str]]]:
        prazo, motivos, _ = detectar_prazo(texto)
        if tipo == "Sentença" and not prazo and "prazo expresso não encontrado" in " ".join(motivos):
            juizado = bool(re.search(r"Juizado|JEC", orgao, re.I))
            recurso = ("Recurso inominado", 10, "Lei 9.099, art. 42") if juizado else \
                ("Apelação", 15, "CPC, art. 1.003, § 5º")
            return [("Embargos de declaração", 5, ["prazo presumido (CPC, art. 1.022/1.023): ler a sentença"]),
                    (recurso[0], recurso[1], [f"prazo recursal presumido ({recurso[2]})"])]
        if prazo:
            return [("Manifestação/providência", prazo, motivos)]
        return [("", None, motivos)]

    def posicao(self, tribunal: str, comarca: str, venc: dt.date | None) -> int:
        def chave(t, c, ativo, v):
            ordem = ORDEM_TRIBUNAIS.index(t) if t in ORDEM_TRIBUNAIS else len(ORDEM_TRIBUNAIS)
            return (ordem, c, 0 if ativo else 1, v or dt.date.max)
        nova = chave(tribunal, comarca, bool(venc), venc)
        com_dados = [l for l in self.prazos.linhas if l.get("Processo")]
        for linha in com_dados:
            atual = chave(str(linha.get("Tribunal") or ""), str(linha.get("Comarca") or ""),
                          str(linha.get("Status") or "") in ATIVOS and bool(para_data(linha.get("Vencimento interno"))),
                          para_data(linha.get("Vencimento interno")))
            if atual > nova:
                return linha["_n"]
        return (max(l["_n"] for l in com_dados) + 1) if com_dados else LINHA_CAB_PRAZOS + 1

    def inserir_prazo(self, valores: dict) -> None:
        n = self.posicao(valores["Tribunal"], valores["Comarca"], valores.get("Vencimento interno"))
        formulas = formulas_prazos(n, self.prazos.letra, self.processos.letra)
        linha = self.linha_completa(self.prazos, valores, formulas)
        self.rel.acoes.append(f"Prazos: inserir na linha {n}: {valores['Processo']} {valores.get('Natureza do prazo', '')}"
                              f" venc. {valores.get('Vencimento interno') or '—'}")
        if not self.ensaio:
            self.p.inserir_linhas("Prazos", n, [linha], herdar_formato_de_cima=n > LINHA_CAB_PRAZOS + 1)
        self.prazos.deslocar(n)
        self.prazos.linhas.append({"_n": n, **valores})
        self.rel.novas += 1

    def registrar_processo(self, numero: str, tribunal: str, comarca: str, orgao: str) -> None:
        d = cnj.digitos(numero)
        if any(cnj.digitos(str(p.get("Processo") or "")) == d for p in self.processos.linhas):
            return
        vazia = next((p for p in self.processos.linhas if not p.get("Processo")), None)
        n = vazia["_n"] if vazia else (max([p["_n"] for p in self.processos.linhas] or [1]) + 1)
        valores = {"Tribunal": tribunal, "Comarca": comarca, "Órgão julgador": orgao, "Processo": cnj.formatar(d),
                   "Instância atual": instancia(orgao), "Situação": "Ativo"}
        formulas = formulas_processos(n, self.processos.letra, self.prazos.letra)
        linha = self.linha_completa(self.processos, valores, formulas)
        self.rel.acoes.append(f"Processos: linha {n}: {cnj.formatar(d)} ({comarca})")
        if not self.ensaio:
            self.p.atualizar("Processos", [{"range": f"A{n}:{get_column_letter(self.processos.ncols)}{n}",
                                            "values": [linha]}])
        if vazia:
            vazia.update(valores)
        else:
            self.processos.linhas.append({"_n": n, **valores})
        self.rel.processos_novos += 1

    def registrar_publicacao(self, valores: dict) -> None:
        d = cnj.digitos(valores["Processo"])
        for p in self.publicacoes.linhas:
            if cnj.digitos(str(p.get("Processo") or "")) != d:
                continue
            disp_existente = para_data(p.get("Disponibilização"))
            if disp_existente == valores["Disponibilização"] and str(p.get("Teor / resumo") or "")[:80] == \
                    valores["Teor / resumo"][:80]:
                return
            if not disp_existente and para_data(p.get("Publicação")) == valores["Publicação"]:
                return  # linha migrada (sem disponibilização) do mesmo dia
        n = max([p["_n"] for p in self.publicacoes.linhas] or [1]) + 1
        linha = self.linha_completa(self.publicacoes, valores, {})
        if not self.ensaio:
            self.p.atualizar("Publicações", [{"range": f"A{n}:{get_column_letter(self.publicacoes.ncols)}{n}",
                                              "values": [linha]}])
        self.publicacoes.linhas.append({"_n": n, **valores})
        self.rel.publicacoes += 1

    def processar_djen(self, itens: list[dict]) -> None:
        self.rel.lidas = len(itens)
        existentes_id = {str(l.get("ID")) for l in self.prazos.linhas if l.get("ID")}
        existentes_pub = {(cnj.digitos(str(l.get("Processo") or "")), para_data(l.get("Publicação")))
                          for l in self.prazos.linhas if l.get("Processo")}
        for item in sorted(itens, key=lambda i: str(i.get("data_disponibilizacao") or "")):
            d = cnj.digitos(str(item.get("numero_processo") or ""))
            disp = para_data(str(item.get("data_disponibilizacao") or "")[:10])
            if len(d) != 20 or not disp:
                self.rel.erros.append(f"item do DJEN sem processo ou data: {str(item)[:80]}")
                continue
            texto = " ".join(str(item.get("texto") or "").split())
            orgao = str(item.get("nomeOrgao") or "")
            trib, cod = cnj.tribunal(d) or "", cnj.comarca(d)
            cal = self.calendario(trib, cod)
            pub = cal.somar_dias_uteis(disp, 1)
            comarca = self.comarca(trib, cod)
            tipo = tipo_ato(texto)
            self.registrar_publicacao({"Publicação": pub, "Disponibilização": disp, "Tribunal": trib,
                                       "Órgão julgador": orgao, "Processo": cnj.formatar(d), "Tipo de ato": tipo,
                                       "Teor / resumo": texto[:1500], "Fonte": "DJEN (robô)"})
            if (d, pub) in existentes_pub:
                continue
            self.registrar_processo(cnj.formatar(d), trib, comarca, orgao)
            criminal = "criminal" in orgao.lower()
            for natureza, dias, motivos in self.prazos_do_ato(texto, tipo, orgao):
                ident = identificador(d, pub, natureza)
                if ident in existentes_id:
                    continue
                conf = list(motivos)
                if motivo := cnj.diagnosticar(d):
                    conf.append(f"CNJ: {motivo}")
                interno = legal = None
                if criminal:
                    conf.append("processo criminal: prazo em dias corridos (CPP, art. 798); não calculado")
                elif dias:
                    r = calcular_prazo(disp, dias, cal, self.hoje)
                    interno, legal = dt.date.fromisoformat(r.vencimento_interno), dt.date.fromisoformat(r.termo_legal)
                    conf += [a for a in r.avisos if "NÃO CONFERIDO" in a or "Sem feriados" in a]
                if tipo == "Intimação de pauta/audiência":
                    conf.append("audiência ou pauta: conferir data e hora")
                self.inserir_prazo({
                    "Tribunal": trib, "Comarca": comarca, "Órgão julgador": orgao, "Processo": cnj.formatar(d),
                    "Disponibilização": disp, "Publicação": pub, "Prazo (d.u.)": dias or "",
                    "Natureza do prazo": natureza, "Vencimento interno": interno, "Termo legal": legal,
                    "Tipo de ato": _tipo_modelo(tipo), "Providência": "", "Status": "A triar",
                    "Instância": instancia(orgao), "Conferência": "; ".join(conf), "Teor / resumo do ato": texto[:1500],
                    "Fonte": "DJEN (robô)", "ID": ident,
                })
                existentes_id.add(ident)
                self.rel.para_revisao += bool(conf)
            existentes_pub.add((d, pub))

    # ---------------------------------------------------------- 3. registro da execução
    def registrar_execucao(self, fonte: str) -> None:
        valores = self.p.ler("Auditoria")
        marco = next((i for i, l in enumerate(valores, 1) if l and str(l[0]).startswith("Registro de execuções")), None)
        if marco is None:
            self.rel.erros.append("aba Auditoria sem o bloco 'Registro de execuções do robô'")
            return
        n = marco + 2
        while n <= len(valores) and valores[n - 1] and str(valores[n - 1][0]).strip():
            n += 1
        linha = [dt.datetime.now().replace(second=0, microsecond=0), fonte, self.rel.lidas,
                 self.rel.novas + self.rel.completadas, self.rel.para_revisao, " | ".join(self.rel.erros)[:500], VERSAO]
        if not self.ensaio:
            self.p.atualizar("Auditoria", [{"range": f"A{n}:G{n}", "values": [[celula(v) for v in linha]]}])


def _indice(letra: str) -> int:
    n = 0
    for ch in letra:
        n = n * 26 + ord(ch) - 64
    return n - 1


def _tipo_modelo(tipo_triagem: str) -> str:
    return {"Intimação de pauta/audiência": "Audiência", "Expedição de alvará/RPV": "Alvará/RPV",
            "Decisão interlocutória": "Decisão", "Outro": "Publicação"}.get(tipo_triagem, tipo_triagem)


# ------------------------------------------------------------------ Google Sheets (gspread)
class PlanilhaGoogle:
    """Acesso à planilha pela API oficial, com a conta de serviço (arquivo JSON)."""

    def __init__(self, id_planilha: str, credencial: str):
        import gspread  # importado aqui para os testes não dependerem da biblioteca
        self._gspread = gspread
        self.doc = gspread.service_account(filename=credencial).open_by_key(id_planilha)
        self._abas = {}

    def _aba(self, nome):
        if nome not in self._abas:
            self._abas[nome] = self.doc.worksheet(nome)
        return self._abas[nome]

    def _tentar(self, funcao, *a, **k):
        for tentativa in range(5):
            try:
                return funcao(*a, **k)
            except self._gspread.exceptions.APIError as e:
                if getattr(e.response, "status_code", 0) not in (429, 500, 502, 503) or tentativa == 4:
                    raise
                time.sleep(2 ** tentativa * 5)

    def ler(self, nome: str) -> list[list]:
        return self._tentar(self._aba(nome).get_values, value_render_option="UNFORMATTED_VALUE",
                            date_time_render_option="SERIAL_NUMBER")

    def atualizar(self, nome: str, dados: list[dict]) -> None:
        self._tentar(self._aba(nome).batch_update, dados, value_input_option="USER_ENTERED")

    def inserir_linhas(self, nome: str, linha: int, valores: list[list], herdar_formato_de_cima: bool) -> None:
        self._tentar(self._aba(nome).insert_rows, valores, row=linha, value_input_option="USER_ENTERED",
                     inherit_from_before=herdar_formato_de_cima)


def main(argv: list[str] | None = None) -> int:
    from automacoes.robos.robo_prazos import FalhaConsulta, consultar_djen
    import requests

    p = argparse.ArgumentParser(description="Atualiza a Planilha Google de prazos a partir do DJEN.")
    p.add_argument("--config")
    p.add_argument("--ensaio", action="store_true", help="mostra o que faria, sem gravar nada")
    p.add_argument("--hoje", type=dt.date.fromisoformat, default=dt.date.today())
    a = p.parse_args(argv)
    configurar_log()
    cfg = carregar_config(a.config)
    g = cfg["planilha_google"]
    robo = Robo(PlanilhaGoogle(g["id"], g["credencial"]), a.hoje, a.ensaio)
    robo.completar_e_recalcular()
    fonte, codigo = "DJEN", 0
    try:
        dias = cfg.getint("djen", "dias_busca", fallback=15)
        itens = consultar_djen(requests.Session(), cfg["djen"]["oab"], cfg["djen"]["uf"],
                               a.hoje - dt.timedelta(days=dias), a.hoje)
        robo.processar_djen(itens)
    except FalhaConsulta as e:
        fonte, codigo = "DJEN indisponível", 1
        robo.rel.erros.append(f"DJEN: {e}")
        log.error("DJEN falhou: %s. Prazos já existentes foram completados/recalculados; nada novo inserido.", e)
    robo.registrar_execucao(fonte)
    r = robo.rel
    for acao in r.acoes:
        log.info(("[ENSAIO] " if a.ensaio else "") + acao)
    log.info("%sLidas %s | novas %s | completadas %s | recalculadas %s | processos novos %s | "
             "publicações %s | para revisão %s | erros %s", "[ENSAIO] " if a.ensaio else "", r.lidas, r.novas,
             r.completadas, r.recalculadas, r.processos_novos, r.publicacoes, r.para_revisao, len(r.erros))
    return codigo


if __name__ == "__main__":
    raise SystemExit(main())
