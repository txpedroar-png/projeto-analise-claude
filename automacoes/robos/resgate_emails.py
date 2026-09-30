"""Resgate de processos novos a partir dos e-mails de push dos tribunais.

Correções em relação à versão anterior:
- a senha de app sai do código: variável de ambiente GMAIL_SENHA_APP;
- a caixa é aberta só para leitura e as mensagens são lidas com BODY.PEEK[]; o
  FETCH (RFC822) anterior marcava como LIDOS todos os e-mails dos últimos 15 dias;
- o tribunal vinha de split(".")[3:5], que devolve "15.0231" e nunca casava com
  "8.15": nenhum processo era cadastrado. Agora usa automacoes.cnj;
- CNJ com dígito verificador errado é descartado e registrado;
- filtro de remetentes (config [gmail] remetentes), para não cadastrar CNJ citado
  em qualquer e-mail;
- lê o corpo em HTML quando não há texto simples e respeita o charset da mensagem;
- mesmas seis abas de tribunal do robô de prazos.

Uso:
    python -m automacoes.robos.resgate_emails
"""
from __future__ import annotations

import argparse
import datetime as dt
import email
import html
import imaplib
import re
from email.header import decode_header, make_header
from pathlib import Path

from openpyxl import load_workbook

from automacoes import cnj
from automacoes.robos.comum import (
    COL_PROC_1, LINHA_INICIAL, aguardar_arquivo, backup, carregar_config, configurar_log, log,
    mapear_processos, segredo, verificar_livre,
)

MESES_IMAP = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def data_imap(d: dt.date) -> str:
    """IMAP exige mês em inglês; strftime('%b') depende do locale do Windows."""
    return f"{d.day:02d}-{MESES_IMAP[d.month - 1]}-{d.year}"


def texto_da_mensagem(msg: email.message.Message) -> str:
    simples, htmls = [], []
    for parte in msg.walk() if msg.is_multipart() else [msg]:
        tipo = parte.get_content_type()
        if tipo not in ("text/plain", "text/html"):
            continue
        bruto = parte.get_payload(decode=True) or b""
        conteudo = bruto.decode(parte.get_content_charset() or "utf-8", errors="replace")
        (simples if tipo == "text/plain" else htmls).append(conteudo)
    if simples:
        return "\n".join(simples)
    return html.unescape(re.sub(r"<[^>]+>", " ", "\n".join(htmls)))


def assunto(msg: email.message.Message) -> str:
    try:
        return str(make_header(decode_header(msg.get("Subject", ""))))
    except (UnicodeDecodeError, LookupError):
        return str(msg.get("Subject", ""))


def remetente_aceito(msg: email.message.Message, remetentes: list[str]) -> bool:
    de = str(msg.get("From", "")).lower()
    return not remetentes or any(r in de for r in remetentes)


def extrair_processos(msg: email.message.Message) -> list[tuple[str, str]]:
    """[(CNJ formatado, resumo)] com os números válidos da mensagem."""
    tit = assunto(msg)
    corpo = texto_da_mensagem(msg)
    resumo = f"{tit} | {' '.join(corpo.split())[:100]}".strip()
    auditoria = cnj.auditar_texto(f"{tit}\n{corpo}")
    for numero, motivo in auditoria["invalidos"].items():
        log.warning("E-mail '%s': CNJ %s descartado (%s)", tit[:60], numero, motivo)
    return [(numero, resumo) for numero in auditoria["bem_formados"]]


def buscar_emails(usuario: str, senha: str, desde: dt.date, remetentes: list[str]) -> list[tuple[str, str]]:
    processos = []
    with imaplib.IMAP4_SSL("imap.gmail.com") as caixa:
        caixa.login(usuario, senha)
        caixa.select("INBOX", readonly=True)
        _, dados = caixa.search(None, f'(SINCE "{data_imap(desde)}")')
        ids = dados[0].split()
        log.info("%s e-mails desde %s", len(ids), desde)
        for i in ids:
            _, partes = caixa.fetch(i, "(BODY.PEEK[])")
            for parte in partes:
                if isinstance(parte, tuple):
                    msg = email.message_from_bytes(parte[1])
                    if remetente_aceito(msg, remetentes):
                        processos.extend(extrair_processos(msg))
    return processos


def primeira_linha_vazia(ws, coluna: int) -> int:
    r = LINHA_INICIAL
    while ws.cell(row=r, column=coluna).value not in (None, ""):
        r += 1
    return r


def cadastrar(wb, processos: list[tuple[str, str]]) -> int:
    existentes = set(mapear_processos(wb))
    novos = 0
    for numero, resumo in processos:
        d = cnj.digitos(numero)
        if d in existentes:
            continue
        aba = cnj.tribunal(numero)
        if aba not in wb.sheetnames:
            log.info("%s: sem aba %s na planilha; não cadastrado", numero, aba)
            continue
        ws = wb[aba]
        linha = primeira_linha_vazia(ws, COL_PROC_1)
        ws.cell(row=linha, column=COL_PROC_1, value=numero)
        ws.cell(row=linha, column=COL_PROC_1 + 1, value=resumo)
        existentes.add(d)
        novos += 1
        log.info("Novo: %s em %s, linha %s", numero, aba, linha)
    return novos


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Cadastra processos novos a partir dos e-mails dos tribunais.")
    p.add_argument("--config")
    a = p.parse_args(argv)
    configurar_log()
    cfg = carregar_config(a.config)
    planilha = Path(cfg["planilha"]["caminho"])
    g = cfg["gmail"]
    remetentes = [r.strip().lower() for r in g.get("remetentes", "").split(",") if r.strip()]
    if not remetentes:
        log.warning("Sem filtro de remetentes: qualquer CNJ citado em qualquer e-mail será cadastrado.")

    desde = dt.date.today() - dt.timedelta(days=g.getint("dias_busca", fallback=15))
    processos = buscar_emails(g["email"], segredo("GMAIL_SENHA_APP"), desde, remetentes)
    if not processos:
        log.info("Nenhum processo encontrado nos e-mails.")
        return 0

    aguardar_arquivo(planilha)
    verificar_livre(planilha)
    backup(planilha, "resgate_emails")
    wb = load_workbook(planilha)
    novos = cadastrar(wb, processos)
    if novos:
        wb.save(planilha)
    log.info("Concluído: %s processos novos cadastrados", novos)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
