"""Robô de gestão processual – Araújo, Azevedo e Costa (v6).

Correções em relação à v5:
- falha do DJEN (403, rede) agora é registrada e aciona o DataJud; na v5 a lista vazia
  impedia o DataJud de rodar e o log dizia "concluída com sucesso";
- tentativas limitadas (a v5 repetia para sempre em erro de rede);
- paginação não depende do campo "totalPages";
- a data na planilha só é substituída por data MAIS RECENTE (na v5, a ordem da API decidia);
- prazo detectado pela mesma regra da triagem (sem confundir horas, meses e anos com dias);
- vencimento, termo legal e criticidade calculados em automacoes.prazos;
- linhas ambíguas vão para revisar_ia.txt, ao lado da planilha;
- hash de movimento do DataJud não quebra com código numérico; aliases TRF e TJDFT corrigidos.

Uso:
    python -m automacoes.robos.robo_prazos
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import re
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

from automacoes import cnj
from automacoes.prazos import Calendario, calcular_prazo
from automacoes.robos.comum import (
    backup, aguardar_arquivo, carregar_config, configurar_log, log, mapear_processos, verificar_livre,
)
from automacoes.triagem import detectar_prazo, tipo_ato

DJEN_URL = "https://comunicaapi.pje.jus.br/api/v1/comunicacao"
DATAJUD_BASE = "https://api-publica.datajud.cnj.jus.br"
# Chave pública divulgada pelo CNJ na wiki do DataJud; o CNJ pode trocá-la.
DATAJUD_CHAVE_PUBLICA = "cDZHYzlZa0JadVREZDJCendQbXY6SkJlTzNjLV9TRENyQk1RdnFKZGRQdw=="
ITENS_POR_PAGINA = 100
MAX_PAGINAS = 50
TENTATIVAS = 3
ABA_TRIAGEM = "Triagem_Diaria"
ABA_CONTROLE_DATAJUD = "Controle_DataJud"
CABECALHO_TRIAGEM = [
    "Data Disponibilização", "Unidade", "Processo", "Alerta", "Prazo (dias)", "Vencimento interno",
    "Termo legal", "Dias úteis restantes", "Criticidade", "Revisar", "Resumo",
]
RE_SEGUNDO_GRAU = re.compile(r"\b(turma|câmara|relator|desembargador)\b", re.I)


class FalhaConsulta(Exception):
    pass


def _get(sessao: requests.Session, url: str, **kw) -> requests.Response:
    ultimo = None
    for tentativa in range(1, TENTATIVAS + 1):
        try:
            resp = sessao.get(url, timeout=25, **kw)
            if resp.status_code == 200:
                return resp
            ultimo = f"HTTP {resp.status_code}: {resp.text[:200]}"
            if resp.status_code in (401, 403, 404):
                break  # recusa definitiva: repetir não resolve
        except requests.RequestException as e:
            ultimo = repr(e)
        time.sleep(5 * tentativa)
    raise FalhaConsulta(ultimo)


def consultar_djen(sessao, oab: str, uf: str, inicio: dt.date, fim: dt.date) -> list[dict]:
    itens, bloco = [], inicio
    while bloco <= fim:
        fim_bloco = min(bloco + dt.timedelta(days=6), fim)
        for pagina in range(1, MAX_PAGINAS + 1):
            params = {
                "numeroOab": oab, "ufOab": uf,
                "dataDisponibilizacaoInicio": bloco.isoformat(), "dataDisponibilizacaoFim": fim_bloco.isoformat(),
                "itensPorPagina": ITENS_POR_PAGINA, "pagina": pagina,
            }
            dados = _get(sessao, DJEN_URL, params=params, headers={"Accept": "application/json"}).json()
            lote = dados.get("items") or []
            itens.extend(lote)
            if len(lote) < ITENS_POR_PAGINA:
                break
        else:
            log.warning("DJEN: bloco %s a %s atingiu %s páginas; resultado pode estar incompleto", bloco, fim_bloco, MAX_PAGINAS)
        bloco = fim_bloco + dt.timedelta(days=1)
    return itens


def alias_datajud(numero: str) -> str | None:
    trib = cnj.tribunal(numero)
    if not trib or not (trib.startswith("TJ") or trib.startswith("TRF")):
        return None
    return f"api_publica_{trib.lower()}"


def consultar_datajud(sessao, numero: str, chave: str) -> dict | None:
    alias = alias_datajud(numero)
    if not alias:
        return None
    corpo = {"query": {"match": {"numeroProcesso": cnj.digitos(numero)}}, "size": 1, "_source": ["movimentos"]}
    resp = sessao.post(f"{DATAJUD_BASE}/{alias}/_search", json=corpo, timeout=20,
                       headers={"Authorization": f"APIKey {chave}"})
    if resp.status_code != 200:
        raise FalhaConsulta(f"DataJud HTTP {resp.status_code}")
    hits = resp.json().get("hits", {}).get("hits", [])
    movs = hits[0]["_source"].get("movimentos", []) if hits else []
    return max(movs, key=lambda m: m.get("dataHora", ""), default=None)


def hash_movimento(mov: dict) -> str:
    return hashlib.md5(f"{mov.get('dataHora', '')}|{mov.get('codigo', '')}".encode()).hexdigest()


def ler_data(valor) -> dt.date | None:
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    s = str(valor or "").strip()
    if re.match(r"\d{4}-\d{2}-\d{2}", s):
        return dt.date.fromisoformat(s[:10])
    if re.match(r"\d{2}/\d{2}/\d{4}", s):
        return dt.datetime.strptime(s[:10], "%d/%m/%Y").date()
    return None


def gravar_data_mais_recente(wb, alvo: dict, data: dt.date) -> bool:
    celula = wb[alvo["aba"]].cell(row=alvo["linha"], column=alvo["col_data"])
    atual = ler_data(celula.value)
    if atual and atual >= data:
        return False
    celula.value = data
    return True


def linha_triagem(data: dt.date | None, unidade: str, numero: str, texto: str, calendarios: dict, hoje: dt.date) -> dict:
    prazo, motivos, _ = detectar_prazo(texto)
    tipo = tipo_ato(texto)
    alerta = f"PRAZO {prazo} DIAS" if prazo else {
        "Intimação de pauta/audiência": "AUDIÊNCIA", "Sentença": "SENTENÇA"}.get(tipo, "Publicação")
    if tipo == "Intimação de pauta/audiência":
        motivos.append("audiência: conferir data/hora e agenda")
    linha = dict.fromkeys(CABECALHO_TRIAGEM, "")
    linha.update({"Data Disponibilização": data, "Unidade": unidade, "Alerta": alerta,
                  "Processo": cnj.formatar(numero) if len(cnj.digitos(numero)) == 20 else numero,
                  "Resumo": " ".join(texto.split())[:300]})
    trib = cnj.tribunal(numero) if not cnj.diagnosticar(numero) else None
    if not trib:
        motivos.append("CNJ inválido ou ausente")
    if prazo and data and trib:
        chave = (trib, cnj.comarca(numero))
        if chave not in calendarios:
            calendarios[chave] = Calendario(tribunal=trib, comarca=chave[1])
        r = calcular_prazo(data, prazo, calendarios[chave], hoje)
        linha.update({"Prazo (dias)": prazo, "Vencimento interno": dt.date.fromisoformat(r.vencimento_interno),
                      "Termo legal": dt.date.fromisoformat(r.termo_legal),
                      "Dias úteis restantes": r.dias_uteis_restantes, "Criticidade": r.criticidade})
        motivos += [a for a in r.avisos if "NÃO CONFERIDO" in a or "Sem feriados" in a]
    linha["Revisar"] = "; ".join(motivos)
    return linha


def trilha_djen(wb, mapa, itens, calendarios, hoje) -> list[dict]:
    triagem = []
    for item in itens:
        numero = str(item.get("numero_processo") or "")
        texto = str(item.get("texto") or "")
        data = ler_data(item.get("data_disponibilizacao"))
        ocorrencias = mapa.get(cnj.digitos(numero))
        if ocorrencias and data:
            alvo = ocorrencias[0]
            if RE_SEGUNDO_GRAU.search(texto):
                alvo = next((o for o in ocorrencias if o["instancia"] == "2ª"), alvo)
            if gravar_data_mais_recente(wb, alvo, data):
                log.info("Data %s gravada: %s em %s linha %s", data, cnj.formatar(numero), alvo["aba"], alvo["linha"])
        triagem.append(linha_triagem(data, item.get("nomeOrgao", "Desconhecida"), numero, texto, calendarios, hoje))
    return triagem


def trilha_datajud(wb, mapa, sessao, chave) -> list[dict]:
    controle = {}
    if ABA_CONTROLE_DATAJUD in wb.sheetnames:
        controle = {str(p).strip(): str(h).strip() for p, h, *_ in
                    wb[ABA_CONTROLE_DATAJUD].iter_rows(min_row=2, values_only=True) if p and h}
    triagem = []
    for numero, ocorrencias in mapa.items():
        try:
            mov = consultar_datajud(sessao, numero, chave)
        except (FalhaConsulta, requests.RequestException) as e:
            log.warning("DataJud %s: %s", cnj.formatar(numero), e)
            continue
        if not mov or hash_movimento(mov) == controle.get(numero):
            continue
        controle[numero] = hash_movimento(mov)
        data = ler_data(mov.get("dataHora"))
        if data:
            gravar_data_mais_recente(wb, ocorrencias[0], data)
        nome = mov.get("nome", "Movimento")
        linha = dict.fromkeys(CABECALHO_TRIAGEM, "")
        linha.update({"Data Disponibilização": data, "Unidade": "DataJud", "Processo": cnj.formatar(numero),
                      "Alerta": f"NOVO MOVIMENTO: {nome}", "Revisar": "movimento do DataJud: ler o teor no PJe",
                      "Resumo": nome})
        triagem.append(linha)
    if ABA_CONTROLE_DATAJUD in wb.sheetnames:
        del wb[ABA_CONTROLE_DATAJUD]
    ws = wb.create_sheet(ABA_CONTROLE_DATAJUD)
    ws.append(["Processo", "Hash_Ultimo_Movimento"])
    for proc, h in controle.items():
        ws.append([proc, h])
    ws.sheet_state = "hidden"
    return triagem


def gravar_triagem(wb, triagem: list[dict]) -> None:
    # Limpa e reaproveita a aba, preservando formatação condicional e posição.
    if ABA_TRIAGEM in wb.sheetnames:
        ws = wb[ABA_TRIAGEM]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(ABA_TRIAGEM)
    ws.append(CABECALHO_TRIAGEM)
    triagem.sort(key=lambda l: (l["Vencimento interno"] or dt.date.max, l["Processo"]))
    for linha in triagem:
        ws.append([linha[c] for c in CABECALHO_TRIAGEM])
    for col in (1, 6, 7):
        for (celula,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            celula.number_format = "DD/MM/YYYY"


def gravar_revisao(pasta: Path, triagem: list[dict]) -> Path | None:
    revisar = [l for l in triagem if l["Revisar"]]
    destino = pasta / "revisar_ia.txt"
    if not revisar:
        destino.unlink(missing_ok=True)
        return None
    blocos = [f"#{i} | {l['Processo']} | disp. {l['Data Disponibilização']} | {l['Revisar']}\n{l['Resumo']}"
              for i, l in enumerate(revisar, 1)]
    destino.write_text(
        "Para cada publicação, responda só: #id | tipo de ato | providência | prazo em dias úteis "
        "(ou 'sem prazo') | fundamento. Não calcule datas. O resumo tem até 300 caracteres: se não bastar, "
        "diga 'ler íntegra'.\n\n" + "\n\n".join(blocos) + "\n", encoding="utf-8")
    return destino


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Robô de prazos: DJEN (com DataJud de reserva) -> planilha.")
    p.add_argument("--config")
    p.add_argument("--hoje", type=dt.date.fromisoformat, default=dt.date.today())
    a = p.parse_args(argv)
    configurar_log()
    cfg = carregar_config(a.config)
    planilha = Path(cfg["planilha"]["caminho"])
    aguardar_arquivo(planilha)
    verificar_livre(planilha)
    backup(planilha, "robo_prazos")

    wb = load_workbook(planilha)
    mapa = mapear_processos(wb)
    log.info("%s processos mapeados nas abas dos tribunais", len(mapa))
    sessao, calendarios, codigo = requests.Session(), {}, 0
    inicio = a.hoje - dt.timedelta(days=cfg.getint("djen", "dias_busca", fallback=180))
    try:
        itens = consultar_djen(sessao, cfg["djen"]["oab"], cfg["djen"]["uf"], inicio, a.hoje)
        log.info("DJEN: %s comunicações", len(itens))
        triagem = trilha_djen(wb, mapa, itens, calendarios, a.hoje)
    except FalhaConsulta as e:
        log.error("DJEN falhou (%s). Acionando DataJud.", e)
        codigo = 1
        triagem = trilha_datajud(wb, mapa, sessao, os.environ.get("DATAJUD_API_KEY") or DATAJUD_CHAVE_PUBLICA)

    gravar_triagem(wb, triagem)
    wb.save(planilha)
    revisao = gravar_revisao(planilha.parent, triagem)
    n_rev = sum(bool(l["Revisar"]) for l in triagem)
    log.info("Concluído: %s linhas na triagem, %s calculadas em Python, %s para revisão%s",
             len(triagem), sum(bool(l["Vencimento interno"]) for l in triagem), n_rev,
             f" ({revisao.name})" if revisao else "")
    return codigo


if __name__ == "__main__":
    raise SystemExit(main())
