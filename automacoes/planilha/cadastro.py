"""Leitura do Cadastro Central de Clientes e vínculo com os processos.

O cadastro é gerado por uma rotina que percorre as pastas dos clientes. Colunas esperadas (aba
"Página1"): Data Extração | Nome do Cliente | CPF | Prioridade | Classificação / Comarca |
Espécie da Demanda | Ação Sugerida | Pendências Probatórias? | Link da Ficha | Link da Pasta.

Todas as abas são varridas, porque a rotina já gravou cliente na aba do painel por engano.
A chave é o CPF (validado pelos dígitos verificadores); nome serve só para sugerir vínculo.
"""
from __future__ import annotations

import datetime as dt
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

RE_CPF = re.compile(r"^\D*(\d{3})\.?(\d{3})\.?(\d{3})-?(\d{2})\D*$")
RE_HYPERLINK = re.compile(r'HYPERLINK\(\s*"([^"]+)"', re.I)
COMARCAS_CONHECIDAS = [
    "João Pessoa", "Mamanguape", "Rio Tinto", "Santa Rita", "Baía da Traição", "Baia da Traição", "Guarabira",
    "Pedras de Fogo", "Campina Grande", "Alagoinha", "Piancó", "Bayeux", "Cabedelo", "Sapé", "Canguaretama",
]
AREAS = {"ações bancárias": "Bancário", "saúde suplementar": "Saúde suplementar", "consumidor": "Consumo geral",
         "civil geral": "Outra"}
ROTULOS = ["data", "nome", "cpf", "prioridade", "classificacao", "especie", "acao", "pendencias", "ficha", "pasta"]


def normalizar(nome: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(nome or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sem_acento).strip().upper()


def cpf_digitos(valor) -> str | None:
    m = RE_CPF.match(str(valor or ""))
    return "".join(m.groups()) if m else None


def cpf_valido(d: str) -> bool:
    if not d or len(d) != 11 or d == d[0] * 11:
        return False
    for n in (9, 10):
        soma = sum(int(d[i]) * (n + 1 - i) for i in range(n))
        if (soma * 10 % 11) % 10 != int(d[n]):
            return False
    return True


def formatar_cpf(d: str) -> str:
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"


def link(celula) -> str:
    if celula.hyperlink is not None and celula.hyperlink.target:
        return celula.hyperlink.target
    v = str(celula.value or "")
    if m := RE_HYPERLINK.search(v):
        return m.group(1)
    return v if v.startswith("http") else ""


def classificar(classificacao: str, especie: str) -> tuple[str, str]:
    """(área, comarca provável) a partir do caminho de pastas."""
    partes = [p.strip() for p in str(classificacao or "").split(">")]
    partes = [p for p in partes if p and not p.lower().startswith("araújo")]
    area = AREAS.get(partes[0].lower(), "") if partes else ""
    if area == "Consumo geral" and re.search(r"aére|LATAM|Azul|GOL\b|bagagem|voo", f"{classificacao} {especie}", re.I):
        area = "Aeronáutico"
    if area == "Consumo geral" and re.search(r"Banc|consignad|RMC|financiamento", especie or "", re.I):
        area = "Bancário"
    texto = normalizar(classificacao)
    comarca = next((c for c in COMARCAS_CONHECIDAS if normalizar(c) in texto), "")
    return area, comarca.replace("Baia", "Baía")


def situacao_documental(v) -> str:
    s = str(v or "").upper()
    if "OK" in s:
        return "Pronto para ajuizamento"
    if "SIM" in s:
        return "Pendente de documentos"
    return ""


def ler_cadastro(caminho: Path) -> tuple[list[dict], list[str]]:
    """(clientes únicos por CPF, avisos)."""
    wb = load_workbook(caminho)
    por_cpf: dict[str, dict] = {}
    avisos: list[str] = []
    for ws in wb:
        for row in ws.iter_rows(min_row=1):
            if len(row) < 10:
                continue
            d = cpf_digitos(row[2].value)
            if not d:
                continue
            nome = str(row[1].value or "").strip()
            conferencia = []
            if not cpf_valido(d):
                conferencia.append("CPF com dígito verificador inválido")
            if ws.title.lower() != "página1":
                conferencia.append(f"linha encontrada fora da base (aba '{ws.title}', linha {row[0].row})")
            area, comarca = classificar(row[4].value, row[5].value)
            entrada = row[0].value
            if isinstance(entrada, str):
                try:
                    entrada = dt.datetime.strptime(entrada.strip()[:16], "%d/%m/%Y %H:%M")
                except ValueError:
                    pass
            registro = {
                "CPF": formatar_cpf(d), "Nome": nome, "Prioridade / hipervulnerabilidade": str(row[3].value or ""),
                "Área": area, "Comarca provável": comarca, "Espécie da demanda": str(row[5].value or ""),
                "Ação sugerida": str(row[6].value or ""), "Situação documental": situacao_documental(row[7].value),
                "Ficha": link(row[8]), "Pasta": link(row[9]), "Data de entrada": entrada,
                "Classificação original": str(row[4].value or ""), "Conferência": "; ".join(conferencia),
            }
            if d in por_cpf:
                anterior = por_cpf[d]
                avisos.append(f"CPF {formatar_cpf(d)} repetido ({anterior['Nome']} / {nome}): mantida a extração mais recente")
                datas = [x for x in (anterior["Data de entrada"], entrada) if isinstance(x, dt.datetime)]
                if datas and entrada != max(datas):
                    continue
                registro["Conferência"] = "; ".join(filter(None, [registro["Conferência"], "CPF repetido no cadastro"]))
            por_cpf[d] = registro
    return list(por_cpf.values()), avisos


def sugerir_vinculo(nome_processo: str, clientes: list[dict]) -> dict | None:
    """Cliente do cadastro com o mesmo nome completo (sem acento/caixa); vínculo final é humano."""
    alvo = normalizar(nome_processo)
    if len(alvo.split()) < 2:
        return None
    iguais = [c for c in clientes if normalizar(c["Nome"]) == alvo]
    return iguais[0] if len(iguais) == 1 else None
