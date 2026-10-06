"""Gera o modelo v2 da planilha de gestão processual e migra os dados da planilha anterior.

Entrada: exportação .xlsx da planilha anterior (abas "Prazos Processuais" e "Triagem_Diaria").
Saída: .xlsx para upload no Google Drive com conversão para Google Sheets.

Princípios:
- uma base por assunto (Prazos, Processos, Decisões, Publicações), sem cópias manuais entre abas;
- datas de prazo calculadas pelo motor testado (automacoes.prazos), não por fórmula;
- fórmulas só para o que muda com o dia (dias restantes, criticidade, contadores);
- colunas da equipe (status, protocolo, observações, link) nunca são escritas pelo robô;
- tudo o que a migração deduziu fica dito na coluna "Conferência".

Uso:
    python -m automacoes.planilha.gerar_modelo origem.xlsx -o modelo_v2.xlsx --hoje 2026-10-06
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from automacoes import cnj
from automacoes.planilha.cadastro import ler_cadastro, normalizar, sugerir_vinculo
from automacoes.prazos import CSV_PADRAO, Calendario, calcular_prazo
from automacoes.triagem import detectar_prazo

# ---------------------------------------------------------------- vocabulário
STATUS = [
    "A triar", "Em elaboração", "Aguardando cliente/documento", "Minuta em revisão",
    "Protocolado / cumprido", "Sem providência (ciência)", "Acompanhar (sem ação do escritório)",
    "A verificar (migração)",
]
STATUS_ENCERRADOS = STATUS[4:7]
TIPOS = [
    "Sentença", "Acórdão", "Decisão", "Despacho", "Intimação", "Expediente/ato ordinatório",
    "Audiência", "Pauta de julgamento", "Alvará/RPV", "Publicação",
]
INSTANCIAS = ["1º grau", "2º grau", "Juizado", "Tribunal superior"]
AREAS = ["Bancário", "Bancário – seguros", "Saúde suplementar", "Aeronáutico", "Consumo geral", "Outra"]
RESULTADOS = [
    "Procedente", "Parcialmente procedente", "Improcedente", "Extinção sem mérito", "Extinção da execução",
    "Indeferimento da inicial", "Homologação de acordo", "Tutela deferida", "Tutela indeferida",
    "ED acolhidos", "ED rejeitados", "Declínio de competência", "Outro",
]
ORDEM_TRIBUNAIS = ["TJPB", "TJRN", "TJPE", "TJCE", "TJSP", "TJPR"]

# Código de origem (4 últimos dígitos do CNJ) -> comarca. "S" = coerente com os órgãos
# julgadores da planilha anterior; "N" = conferir.
COMARCAS = {
    ("TJPB", "0000"): ("TJPB – 2º grau (originário)", "S"),
    ("TJPB", "0001"): ("Campina Grande", "S"),
    ("TJPB", "0231"): ("Mamanguape", "S"),
    ("TJPB", "0261"): ("Piancó", "S"),
    ("TJPB", "0521"): ("Alagoinha", "S"),
    ("TJPB", "0581"): ("Rio Tinto", "S"),
    ("TJPB", "2001"): ("João Pessoa (Capital)", "S"),
    ("TJPB", "2003"): ("João Pessoa (Capital)", "N"),
    ("TJPB", "7701"): ("Núcleo 4.0 – Saúde Suplementar", "S"),
    ("TJRN", "5114"): ("Canguaretama", "S"),
    ("TJRN", "5137"): ("Canguaretama", "N"),
    ("TJPE", "2001"): ("Recife (Capital)", "S"),
    ("TJPE", "4810"): ("Não identificada (cód. 4810)", "N"),
    ("TJSP", "0003"): ("São Paulo – Foro Regional III (Jabaquara)", "S"),
}

# ---------------------------------------------------------------- estilos
AZUL, BRONZE, AZUL2, BRONZE2, CINZA = "1B2E46", "7D593E", "1F497D", "9F856C", "7F7F7F"
BRANCO = Font(color="FFFFFF", bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
QUEBRA = Alignment(vertical="top", wrap_text=True)
DATA = "DD/MM/YYYY"


def preencher(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


# ---------------------------------------------------------------- leitura da origem
RE_DISP = re.compile(r"[Dd]isponibilizad[oa](?: no DJ\S*(?: Eletrônico)?)? em (\d{2}/\d{2}/\d{4})")
RE_DATA = re.compile(r"(\d{2}/\d{2}/\d{4})")
RE_PRAZO_TIPO = re.compile(r"PRAZO\s+(\d+)\s+DIAS", re.I)
ROBOTICO = re.compile(r"^(Publicad|Publicação|Disponibilizado|Intimação|Decisão|Expediente|Sentença|Juntada|"
                      r"Embargos|Audiência|Despacho|Concedida|Prazo)", re.I)


def br(d: dt.date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def data_de(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = RE_DATA.search(str(v or ""))
    return dt.datetime.strptime(m.group(1), "%d/%m/%Y").date() if m else None


def instancia(orgao: str) -> str:
    if re.search(r"Câmara|Gabinete|Gab\.|Turma Recursal|Turma", orgao, re.I):
        return "2º grau"
    if re.search(r"JEC|Juizado", orgao, re.I):
        return "Juizado"
    return "1º grau"


def area_por_parte(parte: str, comarca_cod: str) -> str:
    if comarca_cod == "7701":
        return "Saúde suplementar"
    if re.search(r"Seguros|Auto\b|Vida", parte, re.I):
        return "Bancário – seguros"
    if re.search(r"Banco|Bradesco|C6|Itaú|Santander|Caixa|BMG|Pan\b|Mercantil|Crefisa|Agibank", parte, re.I):
        return "Bancário"
    return ""


def tipo_novo(tipo_antigo: str, texto: str) -> str:
    t = f"{tipo_antigo} {texto}"
    if re.search(r"SENTENÇA|Sentença", t):
        return "Sentença"
    if re.search(r"Pauta", t):
        return "Pauta de julgamento"
    if re.search(r"Audiência", t):
        return "Audiência"
    if re.search(r"Alvará|RPV", t):
        return "Alvará/RPV"
    if re.search(r"Acórdão", t):
        return "Acórdão"
    if re.search(r"Decisão|Tutela|Embargos", t):
        return "Decisão"
    if re.search(r"Despacho", t):
        return "Despacho"
    if re.search(r"Expediente|Ato Ordinatório", t, re.I):
        return "Expediente/ato ordinatório"
    if re.search(r"PRAZO|Intimação", t):
        return "Intimação"
    return "Publicação"


def prazos_da_linha(tipo_antigo: str, prov: str, texto: str, orgao: str) -> list[tuple[str, int | None, str]]:
    """[(natureza, dias, aviso)] a partir do que a planilha anterior registrou."""
    juizado = bool(re.search(r"JEC|Juizado", orgao))
    if re.search(r"\bED\b", prov):
        itens = [("Embargos de declaração", 5, "")]
        if re.search(r"\bRI\b", prov) or (juizado and re.search(r"Apelação|Recurso", prov)):
            itens.append(("Recurso inominado", 10, "prazo recursal presumido (Lei 9.099, art. 42)"))
        elif re.search(r"Apelação", prov):
            itens.append(("Apelação", 15, "prazo recursal presumido (CPC, art. 1.003, § 5º)"))
        elif re.search(r"Agravo", prov):
            itens.append(("Agravo", 15, "prazo recursal presumido (CPC, art. 1.003, § 5º)"))
        return itens
    if prov.startswith("Prazo Apelação"):
        return [("Apelação", 15, "prazo recursal presumido (CPC, art. 1.003, § 5º)")]
    if m := re.search(r"Recurso \((\d+) d\.u\.\)", prov):
        return [("Recurso", int(m.group(1)), "")]
    if m := RE_PRAZO_TIPO.search(tipo_antigo):
        natureza = "Pauta de julgamento" if "Pauta" in texto else "Manifestação/providência"
        aviso = "prazo da pauta presumido em 5 dias: conferir o ato" if "Pauta" in texto else ""
        return [(natureza, int(m.group(1)), aviso)]
    if m := re.search(r"Prazo CPC \((\d+) d\.u\.\)", prov):
        return [("Presumido (CPC, art. 218, § 3º)", int(m.group(1)),
                 "sem prazo expresso no registro anterior: 5 dias presumidos (art. 218, § 3º); ler o ato")]
    dias, _, _ = detectar_prazo(prov)
    if dias:
        return [("Manifestação/providência", dias, "")]
    return [("", None, "")]


def dia_util_anterior(cal: Calendario, d: dt.date) -> dt.date:
    d -= dt.timedelta(days=1)
    while d.weekday() >= 5 or cal.feriado_conferido(d):
        d -= dt.timedelta(days=1)
    return d


def migrar(origem: Path, hoje: dt.date):
    wb = load_workbook(origem)
    ws = wb["Prazos Processuais"]
    brutas = []
    for r in range(6, ws.max_row + 1):
        valores = tuple(ws.cell(r, c).value for c in range(1, 15))
        if valores[0]:
            brutas.append(valores)
    vistos, linhas = set(), []
    for v in brutas:  # duplicatas exatas da planilha anterior
        chave = tuple(str(x) for x in v)
        if chave not in vistos:
            vistos.add(chave)
            linhas.append(v)
    duplicatas = len(brutas) - len(linhas)

    triagem = [r for r in wb["Triagem_Diaria"].iter_rows(min_row=2, max_col=5, values_only=True) if r[2]]
    orgao_por_proc: dict[str, str] = {}
    for v in linhas:
        if v[2] and v[2] != "-":
            orgao_por_proc[cnj.digitos(v[0])] = re.sub(r"\s*\(TJ\w+\)$", "", v[2])
    for t in triagem:
        if t[1] and t[1] != "-":
            orgao_por_proc.setdefault(cnj.digitos(t[2]), re.sub(r"\s*\(TJ\w+\)$", "", t[1]))

    calendarios: dict = {}
    prazos, processos, decisoes = OrderedDict(), OrderedDict(), OrderedDict()
    for proc, partes, orgao, _, pub, tipo_ant, prov, _, _, j_ant, _, st_ant, obs, link in linhas:
        d = cnj.digitos(proc)
        numero = cnj.formatar(d)
        trib = cnj.tribunal(d) or ""
        cod = cnj.comarca(d)
        comarca = COMARCAS.get((trib, cod), (f"Cód. {cod}", "N"))[0]
        orgao = orgao_por_proc.get(d, "") if not orgao or orgao == "-" else re.sub(r"\s*\(TJ\w+\)$", "", orgao)
        prov, obs = str(prov or ""), str(obs or "")
        texto = f"{prov} | {obs}"
        pub = data_de(pub)
        cal = calendarios.setdefault((trib, cod), Calendario(CSV_PADRAO, trib or "TJPB", cod))
        conferencias = []
        if motivo := cnj.diagnosticar(d):
            conferencias.append(f"CNJ: {motivo}")

        disp = None
        if m := RE_DISP.search(texto):
            disp = dt.datetime.strptime(m.group(1), "%d/%m/%Y").date()
        else:
            for t in triagem:
                if cnj.digitos(t[2]) == d and data_de(t[0]) == pub and (m := RE_DISP.search(str(t[4]))):
                    disp = dt.datetime.strptime(m.group(1), "%d/%m/%Y").date()
        if not disp and pub:
            disp = dia_util_anterior(cal, pub)
            conferencias.append("disponibilização deduzida (dia útil anterior à publicação)")

        if isinstance(st_ant, str) and st_ant.startswith("="):
            st_ant = "fórmula antiga (No Prazo/Intempestivo)"
        criminal = "Criminal" in orgao
        if criminal:
            conferencias.append("processo criminal: prazo em dias corridos (CPP, art. 798); não calculado")

        for natureza, dias, aviso in prazos_da_linha(str(tipo_ant or ""), prov, texto, orgao):
            conf = list(conferencias) + ([aviso] if aviso else [])
            interno = legal = None
            if dias and disp and not criminal:
                r = calcular_prazo(disp, dias, cal, hoje)
                interno = dt.date.fromisoformat(r.vencimento_interno)
                legal = dt.date.fromisoformat(r.termo_legal)
                nao_conf = [a.split(" ")[0] for a in r.avisos if "NÃO CONFERIDO" in a]
                if nao_conf:
                    conf.append("feriado(s) não conferido(s) na contagem: " + ", ".join(nao_conf))
                antigo = data_de(j_ant)
                if antigo and antigo != legal and natureza in ("Manifestação/providência", "Embargos de declaração",
                                                               "Presumido (CPC, art. 218, § 3º)", "Recurso",
                                                               "Pauta de julgamento"):
                    conf.append(f"planilha anterior indicava {br(antigo)}")
            elif not dias and data_de(j_ant):
                conf.append(f"planilha anterior indicava vencimento {br(data_de(j_ant))} sem prazo registrado")

            encerrado_antes = st_ant == "Sem Prazo" and re.match(r"(Publicação / Intimação simples|Ciência)", prov)
            if encerrado_antes and not dias:
                status = "Sem providência (ciência)"
            elif legal and legal >= hoje and st_ant in ("Aberto", "Tempestivo", "Vence Hoje", "Vence Hoje (ED)",
                                                           "Pauta Virtual", "fórmula antiga (No Prazo/Intempestivo)"):
                status = "A triar"
            else:
                status = "A verificar (migração)"

            obs_equipe = obs if obs and not ROBOTICO.match(obs) else ""
            teor = " | ".join(x for x in (prov, obs if not obs_equipe else "") if x)
            providencia = "" if re.match(r"(Prazo |Intimação|Decisão no|Despacho no|Expediente no|Publicação / Intimação simples)", prov) else prov
            providencia = re.sub(r"^Manifestação / Providência no prazo de \d+ dias$", "Manifestação (ler o ato)", providencia)
            slug = re.sub(r"\W+", "", (natureza or "ato").lower())[:12]
            ident = f"{d}-{(pub or hoje):%Y%m%d}-{slug}"
            if ident in prazos:
                continue
            prazos[ident] = {
                "Tribunal": trib, "Comarca": comarca, "Órgão julgador": orgao, "Processo": numero,
                "Disponibilização": disp, "Publicação": pub, "Prazo (d.u.)": dias, "Natureza do prazo": natureza,
                "Vencimento interno": interno, "Termo legal": legal, "Tipo de ato": tipo_novo(str(tipo_ant or ""), texto),
                "Providência": providencia, "Status": status, "Observações da equipe": obs_equipe,
                "Link minuta / pasta": link or "", "Instância": instancia(orgao), "Conferência": "; ".join(conf),
                "Teor / resumo do ato": teor, "Fonte": "Migração (planilha anterior)",
                "Status anterior (migração)": st_ant or "", "ID": ident,
            }

        cliente, contraria = "", ""
        if partes and " v. " in str(partes):
            cliente, contraria = [x.strip() for x in str(partes).split(" v. ", 1)]
        p = processos.setdefault(d, {"Tribunal": trib, "Comarca": comarca, "Órgão julgador": orgao,
                                     "Processo": numero, "Cliente": "", "Parte contrária": "", "Área / nicho": "",
                                     "Instância atual": instancia(orgao)})
        if cliente:
            p["Cliente"], p["Parte contrária"] = cliente, contraria
        if orgao:
            p["Órgão julgador"], p["Instância atual"] = orgao, instancia(orgao)
        p["Área / nicho"] = p["Área / nicho"] or area_por_parte(p["Parte contrária"], cod)

        resultado = None
        for padrao, res in ((r"[Pp]rocedência [Pp]arcial|procedência em parte", "Parcialmente procedente"),
                            (r"Improcedência", "Improcedente"), (r"[Jj]ulgado procedente", "Procedente"),
                            (r"Extinção (da )?[Ee]xecução", "Extinção da execução"),
                            (r"Indeferimento da inicial", "Indeferimento da inicial"),
                            (r"Tutela (Antecipada|de Urgência) (deferida|Concedida)|Concedida Tutela", "Tutela deferida"),
                            (r"Embargos de Declaração acolhidos", "ED acolhidos"),
                            (r"[Rr]ejeição (de Embargos de Declaração|ED)", "ED rejeitados"),
                            (r"Incompetência", "Declínio de competência")):
            if re.search(padrao, texto):
                resultado = res
                break
        if resultado:
            decisoes.setdefault(f"{d}-{resultado}", {
                "Processo": numero, "Tribunal": trib, "Comarca": comarca, "Órgão julgador": orgao,
                "Data (publicação)": pub, "Tipo de decisão": "Sentença" if "Sentença" in tipo_novo(str(tipo_ant or ""), texto)
                else "Decisão", "Resultado": resultado,
                "Observação técnica": "", "Fonte": "Migração: resultado lido do resumo; conferir o teor",
            })
    for t in triagem:  # processos que só aparecem na triagem
        d = cnj.digitos(t[2])
        trib, cod = cnj.tribunal(d) or "", cnj.comarca(d)
        processos.setdefault(d, {"Tribunal": trib, "Comarca": COMARCAS.get((trib, cod), (f"Cód. {cod}", "N"))[0],
                                 "Órgão julgador": orgao_por_proc.get(d, ""), "Processo": cnj.formatar(d),
                                 "Cliente": "", "Parte contrária": "", "Área / nicho": area_por_parte("", cod),
                                 "Instância atual": instancia(orgao_por_proc.get(d, ""))})
    publicacoes = []
    for t in triagem:
        disp = None
        if m := RE_DISP.search(str(t[4])):
            disp = dt.datetime.strptime(m.group(1), "%d/%m/%Y").date()
        d = cnj.digitos(t[2])
        publicacoes.append({"Publicação": data_de(t[0]), "Disponibilização": disp, "Tribunal": cnj.tribunal(d) or "",
                            "Órgão julgador": orgao_por_proc.get(d, t[1] if t[1] != "-" else ""),
                            "Processo": cnj.formatar(d), "Tipo de ato": tipo_novo(str(t[3]), str(t[4])),
                            "Teor / resumo": t[4], "Fonte": "Migração (Triagem_Diaria)"})
    return list(prazos.values()), list(processos.values()), list(decisoes.values()), publicacoes, duplicatas


# ---------------------------------------------------------------- montagem
PRAZOS_COLS = [
    # (cabeçalho, largura, bloco)
    ("Tribunal", 8, 0), ("Comarca", 20, 0), ("Órgão julgador", 26, 0), ("Processo", 25, 0), ("Cliente", 22, 0),
    ("Disponibilização", 12, 1), ("Publicação", 12, 1), ("Prazo (d.u.)", 7, 1), ("Natureza do prazo", 18, 1),
    ("Vencimento interno", 12, 1), ("Termo legal", 12, 1), ("Dias úteis restantes", 9, 1), ("Criticidade", 12, 1),
    ("Tipo de ato", 16, 2), ("Providência", 34, 2),
    ("Status", 22, 3), ("Data do protocolo", 12, 3), ("Folga no protocolo (d.u.)", 9, 3),
    ("Observações da equipe", 40, 3), ("Link minuta / pasta", 18, 3),
    ("Área / nicho", 16, 4), ("Instância", 10, 4), ("Conferência", 40, 4), ("Teor / resumo do ato", 50, 4),
    ("Fonte", 16, 4), ("Status anterior (migração)", 16, 4), ("ID", 30, 4), ("Chave da vista (auxiliar)", 10, 4),
]
BLOCOS = [("IDENTIFICAÇÃO", AZUL), ("PRAZOS (interno conservador + termo legal)", BRONZE),
          ("PROVIDÊNCIA", AZUL2), ("GESTÃO DA EQUIPE", BRONZE2), ("CLASSIFICAÇÃO, CONFERÊNCIA E MEMÓRIA", CINZA)]
FERIADOS = "Feriados!$A$2:$A$800"


def col(nome: str, cols=PRAZOS_COLS) -> str:
    return get_column_letter([c[0] for c in cols].index(nome) + 1)


def cabecalho(ws, cols, linha=1, cor=AZUL, altura=36):
    for i, (nome, largura, *_) in enumerate(cols, 1):
        c = ws.cell(linha, i, nome)
        c.font, c.fill, c.alignment = BRANCO, preencher(cor), CENTRO
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha].height = altura


def validacao(ws, opcoes, intervalo):
    dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "Escolha um valor da lista.", "Valor fora da lista", True
    ws.add_data_validation(dv)
    dv.add(intervalo)


def ordem_tribunal(t: str) -> int:
    return ORDEM_TRIBUNAIS.index(t) if t in ORDEM_TRIBUNAIS else len(ORDEM_TRIBUNAIS)


def aba_prazos(wb, prazos, hoje, linhas_extra=25):
    ws = wb.create_sheet("Prazos")
    ws.sheet_properties.tabColor = BRONZE
    for bloco, (titulo, cor) in enumerate(BLOCOS):
        idx = [i for i, c in enumerate(PRAZOS_COLS, 1) if c[2] == bloco]
        ws.merge_cells(start_row=1, start_column=idx[0], end_row=1, end_column=idx[-1])
        c = ws.cell(1, idx[0], titulo)
        c.font, c.fill, c.alignment = BRANCO, preencher(cor), CENTRO
        for i in idx:
            ws.cell(1, i).fill = preencher(cor)
    cabecalho(ws, PRAZOS_COLS, linha=2)
    for i, c in enumerate(PRAZOS_COLS, 1):
        ws.cell(2, i).fill = preencher([AZUL, BRONZE, AZUL2, BRONZE2, CINZA][c[2]])

    ativos = lambda p: p["Status"] in ("A triar",) and p["Vencimento interno"]
    prazos.sort(key=lambda p: (ordem_tribunal(p["Tribunal"]), p["Comarca"], not ativos(p),
                               p["Vencimento interno"] or dt.date.max if ativos(p) else
                               dt.date.max - (p["Publicação"] or dt.date.min)))
    L = {n: col(n) for n, *_ in PRAZOS_COLS}
    total = len(prazos) + linhas_extra
    for r in range(3, 3 + total):
        p = prazos[r - 3] if r - 3 < len(prazos) else {}
        for i, (nome, *_) in enumerate(PRAZOS_COLS, 1):
            if nome in p and p[nome] not in (None, ""):
                ws.cell(r, i, p[nome])
        ws[f"{L['Cliente']}{r}"] = f'=IF({L["Processo"]}{r}="","",IFERROR(""&VLOOKUP({L["Processo"]}{r},Processos!$D:${PC["Cliente"]},2,FALSE),""))'
        ws[f"{L['Área / nicho']}{r}"] = f'=IF({L["Processo"]}{r}="","",IFERROR(""&VLOOKUP({L["Processo"]}{r},Processos!$D:${PC["Área / nicho"]},{PROC_COLS.index(("Área / nicho", 16)) - 2},FALSE),""))'
        v, s = f"{L['Vencimento interno']}{r}", f"{L['Status']}{r}"
        ws[f"{L['Dias úteis restantes']}{r}"] = (
            f'=IF({v}="","",IF({v}=TODAY(),0,IF({v}>TODAY(),NETWORKDAYS(TODAY()+1,{v},{FERIADOS}),'
            f'-NETWORKDAYS({v}+1,TODAY(),{FERIADOS}))))')
        d = f"{L['Dias úteis restantes']}{r}"
        encerrados = ",".join(f'{s}="{x}"' for x in STATUS_ENCERRADOS)
        ws[f"{L['Criticidade']}{r}"] = (
            f'=IF({L["Processo"]}{r}="","",IF(OR({encerrados}),"—",IF({s}="A verificar (migração)","VERIFICAR",'
            f'IF({d}="","SEM PRAZO",IF({d}<0,"VENCIDO",IF({d}=0,"VENCE HOJE",IF({d}<=5,"URGENTE","NORMAL")))))))')
        q, t = f"{L['Data do protocolo']}{r}", f"{L['Termo legal']}{r}"
        ws[f"{L['Folga no protocolo (d.u.)']}{r}"] = (
            f'=IF(OR({q}="",{t}=""),"",IF({q}<={t},NETWORKDAYS({q},{t},{FERIADOS})-1,'
            f'-(NETWORKDAYS({t},{q},{FERIADOS})-1)))')
        # Numera as pendências de cada tribunal na ordem da base: "TJPB#1", "TJPB#2"... (usada pelas Vistas).
        a, m, pr = L["Tribunal"], L["Criticidade"], L["Processo"]
        ws[f"{L['Chave da vista (auxiliar)']}{r}"] = (
            f'=IF(OR({pr}{r}="",{m}{r}="—"),"",{a}{r}&"#"&COUNTIFS({a}$3:{a}{r},{a}{r},{m}$3:{m}{r},"<>—",'
            f'{pr}$3:{pr}{r},"<>"))')
        for nome in ("Disponibilização", "Publicação", "Vencimento interno", "Termo legal", "Data do protocolo"):
            ws[f"{L[nome]}{r}"].number_format = DATA
        for nome in ("Providência", "Observações da equipe", "Conferência", "Teor / resumo do ato"):
            ws[f"{L[nome]}{r}"].alignment = QUEBRA
    ultima = 2 + total
    validacao(ws, STATUS, f"{L['Status']}3:{L['Status']}{ultima}")
    validacao(ws, TIPOS, f"{L['Tipo de ato']}3:{L['Tipo de ato']}{ultima}")
    validacao(ws, INSTANCIAS, f"{L['Instância']}3:{L['Instância']}{ultima}")

    tudo = f"A3:{get_column_letter(len(PRAZOS_COLS))}{ultima}"
    encerrados = ",".join(f'${L["Status"]}3="{x}"' for x in STATUS_ENCERRADOS)
    ws.conditional_formatting.add(tudo, FormulaRule(formula=[f"OR({encerrados})"], stopIfTrue=True,
                                                    font=Font(color="A6A6A6", strike=True), fill=preencher("F3F3F3")))
    crit = f"{L['Criticidade']}3:{L['Criticidade']}{ultima}"
    for valor, cor, fonte in (("VENCIDO", "F4CCCC", "990000"), ("VENCE HOJE", "F4CCCC", "990000"),
                              ("URGENTE", "FFF2CC", "7F6000"), ("NORMAL", "D9EAD3", "274E13"),
                              ("VERIFICAR", "FCE5CD", "783F04")):
        ws.conditional_formatting.add(crit, FormulaRule(formula=[f'${L["Criticidade"]}3="{valor}"'],
                                                        fill=preencher(cor), font=Font(color=fonte, bold=True)))
    conf = f"{L['Conferência']}3:{L['Conferência']}{ultima}"
    ws.conditional_formatting.add(conf, FormulaRule(formula=[f'${L["Conferência"]}3<>""'], font=Font(color="B45F06")))
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(PRAZOS_COLS))}{ultima}"
    return ws


PROC_COLS = [
    ("Tribunal", 8), ("Comarca", 20), ("Órgão julgador", 26), ("Processo", 25), ("Cliente", 24),
    ("CPF do cliente", 15), ("Parte contrária", 22), ("Área / nicho", 16), ("Tese principal", 30), ("Hipervulnerabilidade", 22),
    ("Instância atual", 10), ("Fase", 16), ("Situação", 12), ("Valor da causa (R$)", 14),
    ("Pasta do caso (link)", 18), ("Observações", 36), ("Prazos registrados", 9), ("Próximo vencimento interno", 12),
    ("Vínculo com o cadastro", 24), ("Sugestão de vínculo (conferir)", 34),
]
CLI_COLS = [
    ("CPF", 15), ("Nome", 30), ("Prioridade / hipervulnerabilidade", 28), ("Área", 16), ("Comarca provável", 16),
    ("Espécie da demanda", 40), ("Ação sugerida", 40), ("Situação documental", 20), ("Etapa", 22),
    ("Processos vinculados", 10), ("Ficha", 12), ("Pasta", 12), ("Data de entrada", 16), ("Observações", 36),
    ("Conferência", 36), ("Classificação original", 40),
]
PC = {n: get_column_letter(i) for i, (n, _) in enumerate(PROC_COLS, 1)}
CC = {n: get_column_letter(i) for i, (n, _) in enumerate(CLI_COLS, 1)}


def aba_processos(wb, processos, clientes=(), linhas_extra=20):
    ws = wb.create_sheet("Processos")
    ws.sheet_properties.tabColor = AZUL
    cabecalho(ws, PROC_COLS)
    processos.sort(key=lambda p: (ordem_tribunal(p["Tribunal"]), p["Comarca"], p["Processo"]))
    pc = PC
    ultima = 1 + len(processos) + linhas_extra
    for p in processos:
        if p.get("Cliente") and (c := sugerir_vinculo(p["Cliente"], list(clientes))):
            p["Sugestão de vínculo (conferir)"] = f"{c['Nome']} – CPF {c['CPF']} (mesmo nome no cadastro)"
    for r in range(2, ultima + 1):
        p = processos[r - 2] if r - 2 < len(processos) else {}
        for i, (nome, _) in enumerate(PROC_COLS, 1):
            if p.get(nome):
                ws.cell(r, i, p[nome])
        if p:
            ws.cell(r, PROC_COLS.index(("Situação", 12)) + 1, "Ativo")
        d = f"{pc['Processo']}{r}"
        ws[f"{pc['Prazos registrados']}{r}"] = f'=IF({d}="","",COUNTIF(Prazos!$D:$D,{d}))'
        ws[f"{pc['Próximo vencimento interno']}{r}"] = (
            f'=IF({d}="","",IFERROR(1/(1/_xlfn.MINIFS(Prazos!$J:$J,Prazos!$D:$D,{d},Prazos!$J:$J,">="&TODAY())),""))')
        ws[f"{pc['Próximo vencimento interno']}{r}"].number_format = DATA
        ws[f"{pc['Valor da causa (R$)']}{r}"].number_format = '"R$" #,##0.00'
        cpf = f"{pc['CPF do cliente']}{r}"
        ws[f"{pc['Vínculo com o cadastro']}{r}"] = (
            f'=IF({d}="","",IF({cpf}="","Sem CPF: vincular ao cadastro",'
            f'IF(COUNTIF(Clientes!$A:$A,{cpf})=0,"CPF fora do cadastro","Vinculado")))')
    validacao(ws, AREAS, f"{pc['Área / nicho']}2:{pc['Área / nicho']}{ultima}")
    validacao(ws, INSTANCIAS, f"{pc['Instância atual']}2:{pc['Instância atual']}{ultima}")
    validacao(ws, ["Postulatória", "Instrutória", "Sentenciado", "Recursal", "Cumprimento de sentença",
                   "Levantamento de valores", "Arquivado"], f"{pc['Fase']}2:{pc['Fase']}{ultima}")
    validacao(ws, ["Ativo", "Suspenso", "Encerrado"], f"{pc['Situação']}2:{pc['Situação']}{ultima}")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PROC_COLS))}{ultima}"
    vinc = f"{pc['Vínculo com o cadastro']}2:{pc['Vínculo com o cadastro']}{ultima}"
    for valor, cor in (("Vinculado", "D9EAD3"), ("CPF fora do cadastro", "F4CCCC"), ("Sem CPF: vincular ao cadastro", "FCE5CD")):
        ws.conditional_formatting.add(vinc, FormulaRule(formula=[f'${pc["Vínculo com o cadastro"]}2="{valor}"'],
                                                        fill=preencher(cor)))


ETAPAS = ["Pendente de documentos", "Pronto para ajuizar", "Ajuizado"]


def aba_clientes(wb, clientes, linhas_extra=30):
    """Cadastro único de clientes (CPF = chave). Alimentado pelo Cadastro Central; etapa e vínculo por fórmula."""
    ws = wb.create_sheet("Clientes")
    ws.sheet_properties.tabColor = BRONZE
    cabecalho(ws, CLI_COLS, cor=BRONZE)
    clientes = sorted(clientes, key=lambda c: (c["Área"], c["Comarca provável"], normalizar(c["Nome"])))
    ultima = 1 + len(clientes) + linhas_extra
    for r in range(2, ultima + 1):
        c = clientes[r - 2] if r - 2 < len(clientes) else {}
        for i, (nome, _) in enumerate(CLI_COLS, 1):
            if c.get(nome) and nome not in ("Ficha", "Pasta"):
                ws.cell(r, i, c[nome])
        for nome, rotulo in (("Ficha", "Abrir ficha"), ("Pasta", "Abrir pasta")):
            if c.get(nome):
                ws[f"{CC[nome]}{r}"] = f'=HYPERLINK("{c[nome]}","{rotulo}")'
        a = f"{CC['CPF']}{r}"
        ws[f"{CC['Processos vinculados']}{r}"] = f'=IF({a}="","",COUNTIF(Processos!${PC["CPF do cliente"]}:${PC["CPF do cliente"]},{a}))'
        n, sd = f"{CC['Processos vinculados']}{r}", f"{CC['Situação documental']}{r}"
        ws[f"{CC['Etapa']}{r}"] = (f'=IF({a}="","",IF({n}>0,"Ajuizado",IF({sd}="Pronto para ajuizamento",'
                                   f'"Pronto para ajuizar","Pendente de documentos")))')
        ws[f"{CC['Data de entrada']}{r}"].number_format = "DD/MM/YYYY HH:MM"
        for nome in ("Espécie da demanda", "Ação sugerida", "Observações", "Conferência"):
            ws[f"{CC[nome]}{r}"].alignment = QUEBRA
    validacao(ws, AREAS, f"{CC['Área']}2:{CC['Área']}{ultima}")
    validacao(ws, ["Pendente de documentos", "Pronto para ajuizamento"],
              f"{CC['Situação documental']}2:{CC['Situação documental']}{ultima}")
    etapa = f"{CC['Etapa']}2:{CC['Etapa']}{ultima}"
    for valor, cor in zip(ETAPAS, ("FCE5CD", "FFF2CC", "D9EAD3")):
        ws.conditional_formatting.add(etapa, FormulaRule(formula=[f'${CC["Etapa"]}2="{valor}"'], fill=preencher(cor)))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CLI_COLS))}{ultima}"


DEC_COLS = [
    ("Processo", 25), ("Tribunal", 8), ("Comarca", 20), ("Órgão julgador", 26), ("Data (publicação)", 12),
    ("Tipo de decisão", 16), ("Resultado", 22), ("Tese(s) apreciada(s)", 30), ("Repetição em dobro", 12),
    ("Dano moral (R$)", 12), ("Honorários (%)", 10), ("Fundamentação genérica (art. 489, § 1º)?", 14),
    ("Recurso", 14), ("Resultado do recurso", 18), ("Link do inteiro teor", 18), ("Observação técnica", 36),
    ("Fonte", 26),
]


def aba_decisoes(wb, decisoes, linhas_extra=30):
    ws = wb.create_sheet("Decisões")
    ws.sheet_properties.tabColor = AZUL2
    cabecalho(ws, DEC_COLS, cor=AZUL2)
    dc = {n: get_column_letter(i) for i, (n, _) in enumerate(DEC_COLS, 1)}
    ultima = 1 + len(decisoes) + linhas_extra
    for r, d in enumerate(decisoes, 2):
        for i, (nome, _) in enumerate(DEC_COLS, 1):
            if d.get(nome):
                ws.cell(r, i, d[nome])
    for r in range(2, ultima + 1):
        ws[f"{dc['Data (publicação)']}{r}"].number_format = DATA
        ws[f"{dc['Dano moral (R$)']}{r}"].number_format = '"R$" #,##0.00'
    validacao(ws, ["Sentença", "Acórdão", "Decisão", "Decisão monocrática (relator)"],
              f"{dc['Tipo de decisão']}2:{dc['Tipo de decisão']}{ultima}")
    validacao(ws, RESULTADOS, f"{dc['Resultado']}2:{dc['Resultado']}{ultima}")
    validacao(ws, ["Sim", "Não", "Não pedida"], f"{dc['Repetição em dobro']}2:{dc['Repetição em dobro']}{ultima}")
    validacao(ws, ["Sim", "Não"], f"{dc['Fundamentação genérica (art. 489, § 1º)?']}2:"
                                  f"{dc['Fundamentação genérica (art. 489, § 1º)?']}{ultima}")
    validacao(ws, ["Interposto pelo escritório", "Interposto pela parte contrária", "Ambos", "Não houve",
                   "Prazo em curso"], f"{dc['Recurso']}2:{dc['Recurso']}{ultima}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DEC_COLS))}{ultima}"


PUB_COLS = [("Publicação", 12), ("Disponibilização", 12), ("Tribunal", 8), ("Órgão julgador", 28), ("Processo", 25),
            ("Tipo de ato", 18), ("Teor / resumo", 70), ("Fonte", 22)]


def aba_publicacoes(wb, publicacoes):
    ws = wb.create_sheet("Publicações")
    ws.sheet_properties.tabColor = CINZA
    cabecalho(ws, PUB_COLS, cor=CINZA)
    publicacoes.sort(key=lambda p: p["Publicação"] or dt.date.min, reverse=True)
    for r, p in enumerate(publicacoes, 2):
        for i, (nome, _) in enumerate(PUB_COLS, 1):
            if p.get(nome) not in (None, ""):
                ws.cell(r, i, p[nome])
        ws.cell(r, 1).number_format = ws.cell(r, 2).number_format = DATA
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(2, len(publicacoes) + 1)}"


def aba_feriados(wb):
    ws = wb.create_sheet("Feriados")
    ws.sheet_properties.tabColor = CINZA
    cols = [("Data", 12), ("Descrição", 40), ("Abrangência", 12), ("Tipo", 18), ("Conferido", 10), ("Fonte", 60)]
    cabecalho(ws, cols, cor=CINZA)
    linhas = list(csv.DictReader(CSV_PADRAO.open(encoding="utf-8")))
    for ano in (2025, 2026):
        d = dt.date(ano, 12, 20)
        while d <= dt.date(ano + 1, 1, 20):
            linhas.append({"data": d.isoformat(), "descricao": "Suspensão de prazos (recesso)", "abrangencia": "NACIONAL",
                           "tipo": "RECESSO", "conferido": "S", "fonte": "CPC, art. 220"})
            d += dt.timedelta(days=1)
    linhas.sort(key=lambda l: (l["data"], l["abrangencia"]))
    for r, l in enumerate(linhas, 2):
        ws.cell(r, 1, dt.date.fromisoformat(l["data"])).number_format = DATA
        for i, k in enumerate(("descricao", "abrangencia", "tipo", "conferido", "fonte"), 2):
            ws.cell(r, i, l[k])
    validacao(ws, ["S", "N"], "E2:E800")
    ws.conditional_formatting.add("A2:F800", FormulaRule(formula=['$E2="N"'], font=Font(color="B45F06")))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(linhas) + 1}"


def aba_comarcas(wb):
    ws = wb.create_sheet("Comarcas")
    ws.sheet_properties.tabColor = CINZA
    cabecalho(ws, [("Tribunal", 9), ("Código de origem", 10), ("Comarca", 40), ("Conferido", 10),
                   ("Observação", 50)], cor=CINZA)
    for r, ((t, c), (nome, conf)) in enumerate(sorted(COMARCAS.items(), key=lambda x: (ordem_tribunal(x[0][0]), x[0][1])), 2):
        ws.cell(r, 1, t)
        ws.cell(r, 2, c)
        ws.cell(r, 3, nome)
        ws.cell(r, 4, conf)
        ws.cell(r, 5, "Deduzido dos órgãos julgadores registrados" if conf == "S" else
                "Conferir: código não confirmado pelos órgãos registrados")


def aba_vista(wb, tribunal, linhas=200):
    """Pendências do tribunal na ordem da base (comarca, vencimento), por ÍNDICE/CORRESP: funciona no
    Excel, no LibreOffice e no Google Sheets, sem funções de matriz dinâmica."""
    ws = wb.create_sheet(f"Vista {tribunal}")
    ws.sheet_properties.tabColor = AZUL
    ws["A1"] = (f"{tribunal}: prazos e atos pendentes, por comarca e vencimento interno. Somente leitura: "
                "edite na aba Prazos (esta vista se atualiza sozinha).")
    ws["A1"].font = Font(bold=True, color=AZUL)
    ws.merge_cells("A1:P1")
    ultima_col = PRAZOS_COLS.index(next(c for c in PRAZOS_COLS if c[0] == "Status")) + 1
    cabecalho(ws, [("Linha na base", 6)] + [(n, w) for n, w, _ in PRAZOS_COLS[1:ultima_col]], linha=2)
    chave = col("Chave da vista (auxiliar)")
    datas = {col(n) for n in ("Disponibilização", "Publicação", "Vencimento interno", "Termo legal")}
    for r in range(3, 3 + linhas):
        ws[f"A{r}"] = f'=IFERROR(MATCH("{tribunal}#"&(ROW()-2),Prazos!${chave}:${chave},0),"")'
        for i in range(2, ultima_col + 1):
            x = get_column_letter(i)
            ws[f"{x}{r}"] = f'=IF($A{r}="","",IF(INDEX(Prazos!{x}:{x},$A{r})="","",INDEX(Prazos!{x}:{x},$A{r})))'
            if x in datas:
                ws[f"{x}{r}"].number_format = DATA
    crit = col("Criticidade")
    for valor, cor, fonte in (("VENCIDO", "F4CCCC", "990000"), ("VENCE HOJE", "F4CCCC", "990000"),
                              ("URGENTE", "FFF2CC", "7F6000"), ("NORMAL", "D9EAD3", "274E13"),
                              ("VERIFICAR", "FCE5CD", "783F04")):
        ws.conditional_formatting.add(f"{crit}3:{crit}{2 + linhas}", FormulaRule(
            formula=[f'${crit}3="{valor}"'], fill=preencher(cor), font=Font(color=fonte, bold=True)))
    ws.freeze_panes = "E3"


def aba_painel(wb, combinacoes, hoje):
    ws = wb.create_sheet("Painel", 0)
    ws.sheet_properties.tabColor = AZUL
    ws.column_dimensions["A"].width = 30
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 16
    ws["A1"] = "Painel de gestão processual – Araújo, Azevedo e Costa Advogados"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = f"Modelo v2 gerado em {hoje:%d/%m/%Y}. Contadores atualizados automaticamente."
    ws["A2"].font = Font(italic=True, color=CINZA)
    S, C, J = f"Prazos!${col('Status')}:${col('Status')}", f"Prazos!${col('Criticidade')}:${col('Criticidade')}", \
        f"Prazos!${col('Tipo de ato')}:${col('Tipo de ato')}"
    kpis = [
        ("Vencidos sem baixa", f'=COUNTIF({C},"VENCIDO")', "F4CCCC"),
        ("Vencem hoje", f'=COUNTIF({C},"VENCE HOJE")', "F4CCCC"),
        ("Urgentes (1 a 5 dias úteis)", f'=COUNTIF({C},"URGENTE")', "FFF2CC"),
        ("No prazo (mais de 5 dias úteis)", f'=COUNTIF({C},"NORMAL")', "D9EAD3"),
        ("A triar", f'=COUNTIF({S},"A triar")', "FFFFFF"),
        ("Aguardando cliente/documento", f'=COUNTIF({S},"Aguardando cliente/documento")', "FFFFFF"),
        ("Minuta em revisão", f'=COUNTIF({S},"Minuta em revisão")', "FFFFFF"),
        ("A verificar (herdados da migração)", f'=COUNTIF({S},"A verificar (migração)")', "FCE5CD"),
        ("Audiências e pautas pendentes", f'=COUNTIFS({J},"Audiência",{S},"<>Protocolado / cumprido")+'
                                          f'COUNTIFS({J},"Pauta de julgamento",{S},"<>Protocolado / cumprido")', "FFFFFF"),
        ("Clientes no cadastro", f'=COUNTA(Clientes!$A:$A)-1', "FFFFFF"),
        ("Clientes prontos para ajuizar", f'=COUNTIF(Clientes!${CC["Etapa"]}:${CC["Etapa"]},"Pronto para ajuizar")', "FFF2CC"),
        ("Clientes pendentes de documentos",
         f'=COUNTIF(Clientes!${CC["Etapa"]}:${CC["Etapa"]},"Pendente de documentos")', "FCE5CD"),
        ("Processos sem cliente vinculado",
         f'=COUNTIF(Processos!${PC["Vínculo com o cadastro"]}:${PC["Vínculo com o cadastro"]},"Sem CPF*")', "FCE5CD"),
    ]
    ws["A4"], ws["B4"] = "Indicador", "Quantidade"
    for c in ("A4", "B4"):
        ws[c].font, ws[c].fill = BRANCO, preencher(AZUL)
    for r, (rotulo, formula, cor) in enumerate(kpis, 5):
        ws[f"A{r}"], ws[f"B{r}"] = rotulo, formula
        ws[f"A{r}"].fill = ws[f"B{r}"].fill = preencher(cor)
        ws[f"B{r}"].font = Font(bold=True, size=12)
    base = 5 + len(kpis) + 1
    cab = ["Tribunal", "Comarca", "Pendentes", "Urgentes / hoje", "Vencidos", "A verificar", "Próximo vencimento"]
    for i, t in enumerate(cab, 1):
        c = ws.cell(base, i, t)
        c.font, c.fill, c.alignment = BRANCO, preencher(BRONZE), CENTRO
    A, B = "Prazos!$A:$A", "Prazos!$B:$B"
    Jv = f"Prazos!${col('Vencimento interno')}:${col('Vencimento interno')}"
    for r, (trib, comarca) in enumerate(combinacoes, base + 1):
        ws.cell(r, 1, trib)
        ws.cell(r, 2, comarca)
        filtro = f'{A},$A{r},{B},$B{r}'
        ws.cell(r, 3, f'=COUNTIFS({filtro},{C},"<>—",{C},"<>VERIFICAR")')
        ws.cell(r, 4, f'=COUNTIFS({filtro},{C},"URGENTE")+COUNTIFS({filtro},{C},"VENCE HOJE")')
        ws.cell(r, 5, f'=COUNTIFS({filtro},{C},"VENCIDO")')
        ws.cell(r, 6, f'=COUNTIFS({filtro},{C},"VERIFICAR")')
        ws.cell(r, 7, f'=IFERROR(1/(1/_xlfn.MINIFS({Jv},{A},$A{r},{B},$B{r},{Jv},">="&TODAY())),"")').number_format = DATA
    r = base + len(combinacoes) + 2
    ws.cell(r, 1, "Acesso aos sistemas").font = Font(bold=True, color=AZUL)
    portais = [("TJPB – PJe", "https://www.tjpb.jus.br/pje"), ("TJRN – PJe", "https://www.tjrn.jus.br"),
               ("TJPE – PJe", "https://pje.tjpe.jus.br"), ("TJCE – PJe", "https://pje.tjce.jus.br"),
               ("TJSP – e-SAJ", "https://esaj.tjsp.jus.br"), ("Comunica PJe / DJEN", "https://comunica.pje.jus.br")]
    for i, (nome, url) in enumerate(portais, r + 1):
        ws.cell(i, 1, f'=HYPERLINK("{url}","{nome}")')
    ws.freeze_panes = "A4"


LEIA_ME = [
    ("COMO ESTA PLANILHA FUNCIONA", None),
    ("Finalidade", "Uma ferramenta só para prazos, processos, decisões e auditoria. O robô e as fórmulas fazem o trabalho "
                   "mecânico; a equipe registra o que só humanos sabem; a IA é usada para ler atos ambíguos e para a "
                   "auditoria mensal."),
    ("Abas", "Painel (indicadores e mapa por tribunal/comarca) · Prazos (base de trabalho diária) · Vistas por tribunal "
             "(somente leitura) · Processos (cadastro único: cliente, parte, área, tese) · Decisões (memória para "
             "jurimetria) · Publicações (registro bruto) · Feriados · Comarcas · Auditoria · Leia-me."),
    ("DUPLA CONTAGEM DE PRAZOS", None),
    ("Vencimento interno", "Meta do escritório. Conta do 1º dia útil após a DISPONIBILIZAÇÃO; fica 1 dia útil antes do "
                           "termo legal, de propósito (gordura)."),
    ("Termo legal", "Último dia do prazo pelo CPC (publicação = 1º dia útil após a disponibilização; contagem a partir do "
                    "dia útil seguinte). Usar para aferir tempestividade, nunca como meta."),
    ("Quem calcula", "As datas vêm do motor em Python (repositório projeto-analise-claude, automacoes/prazos.py), com "
                     "testes automatizados. Não digite fórmulas nessas colunas. Feriado só adia prazo se estiver "
                     "conferido (S) na aba Feriados; os não conferidos geram aviso na coluna Conferência."),
    ("Dias úteis restantes", "Fórmula viva: conta até o vencimento interno com TODOS os feriados da aba (de qualquer "
                             "tribunal) e o recesso. Pode mostrar 1 dia a menos que o real: erra para o lado seguro."),
    ("Fora do cálculo", "Prazos em dobro, processos criminais (dias corridos, CPP art. 798), prazos em horas ou meses: "
                        "aparecem em Conferência para análise humana."),
    ("O QUE CADA UM PREENCHE", None),
    ("Robô / migração", "Tribunal, comarca, órgão, processo, datas, prazo, natureza, vencimentos, tipo de ato, teor, "
                        "fonte, ID, conferência."),
    ("Equipe", "Status, data do protocolo, observações (formato: dd/mm – iniciais: texto, mais recente primeiro), link "
               "da minuta, providência (quando o robô não souber) e o cadastro da aba Processos (cliente, parte "
               "contrária, área, tese, hipervulnerabilidade). Cliente e área aparecem sozinhos na aba Prazos."),
    ("Clientes", "Cadastro único, com CPF como chave, alimentado pelo Cadastro Central de Clientes. Etapa: Pendente de "
                 "documentos → Pronto para ajuizar → Ajuizado (automático quando algum processo recebe o CPF). Ao "
                 "ajuizar, preencher o CPF do cliente na aba Processos: o vínculo passa a 'Vinculado'. Sugestões de "
                 "vínculo por nome só aparecem com nome completo idêntico e sem homônimos, e sempre exigem conferência."),
    ("Ao protocolar", "Status = Protocolado / cumprido e preencher a Data do protocolo: a linha esmaece e a coluna "
                      "Folga mede quantos dias úteis sobraram até o termo legal (indicador de gestão)."),
    ("Decisões", "Toda sentença, acórdão ou tutela relevante ganha uma linha em Decisões. É a base para saber, com "
                 "dados, como cada comarca e órgão decide cada tese."),
    ("STATUS", None),
] + [(s, d) for s, d in (
    ("A triar", "Ato novo ainda não lido por advogado."),
    ("Em elaboração", "Peça ou providência em andamento."),
    ("Aguardando cliente/documento", "Depende de documento ou informação do cliente."),
    ("Minuta em revisão", "Minuta pronta aguardando revisão (validação humana obrigatória)."),
    ("Protocolado / cumprido", "Cumprido; preencher a data do protocolo."),
    ("Sem providência (ciência)", "Ato que não exige manifestação."),
    ("Acompanhar (sem ação do escritório)", "Prazo da parte contrária, pauta, diligência do juízo."),
    ("A verificar (migração)", "Herdado da planilha anterior sem confirmação de cumprimento: conferir no PJe e "
                               "dar baixa ou registrar a perda de prazo."),
)] + [
    ("REGRAS DO ESCRITÓRIO", None),
    ("Validação humana", "Nenhum prazo, cálculo ou peça vai a protocolo sem revisão de advogado."),
    ("Jurisprudência", "Precedente só entra em peça depois de conferido no repositório oficial."),
    ("LGPD", "Planilha restrita à equipe. Não compartilhar por link aberto; não colar dados de clientes em IA fora das "
             "contas do escritório."),
]


def aba_leia_me(wb):
    ws = wb.create_sheet("Leia-me")
    ws.sheet_properties.tabColor = BRONZE2
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    for r, (a, b) in enumerate(LEIA_ME, 1):
        ws.cell(r, 1, a)
        if b is None:
            ws.cell(r, 1).font, ws.cell(r, 1).fill = BRANCO, preencher(AZUL)
            ws.cell(r, 2).fill = preencher(AZUL)
        else:
            ws.cell(r, 1).font = Font(bold=True)
            ws.cell(r, 2, b).alignment = QUEBRA


AUDITORIA_PLANO = [
    ("Diária (robô)", "Robô roda e registra a execução; equipe lê a aba Prazos filtrando Status = A triar.",
     "Automática + 10 min humanos"),
    ("Semanal (humana)", "Sortear 5 prazos calculados e refazer a contagem à mão no calendário do tribunal; conferir "
                         "duplicidades; nenhuma linha 'A triar' com mais de 2 dias úteis; Conferência sem avisos "
                         "esquecidos.", "15 min"),
    ("Mensal (IA)", "Exportar Prazos e Decisões em CSV e pedir à IA: (1) inconsistências entre tipo de ato, prazo e "
                    "natureza; (2) providências vagas; (3) padrões nas decisões por comarca/órgão/tese. A IA aponta; "
                    "advogado decide.", "Uma conversa"),
    ("Trimestral (crítica)", "Indicadores: folga média no protocolo, vencidos, tempo até a triagem, resultados por "
                             "tese e comarca. Decidir ajustes no modelo, no robô e nas teses.", "1 reunião"),
    ("Anual (dezembro)", "Conferir o calendário de feriados do ano seguinte de cada tribunal (portarias) e marcar "
                         "Conferido = S com a fonte.", "1 hora"),
]


def aba_auditoria(wb):
    ws = wb.create_sheet("Auditoria")
    ws.sheet_properties.tabColor = BRONZE2
    cabecalho(ws, [("Periodicidade", 20), ("O que fazer", 90), ("Esforço", 22)], cor=BRONZE2)
    for r, linha in enumerate(AUDITORIA_PLANO, 2):
        for i, v in enumerate(linha, 1):
            ws.cell(r, i, v).alignment = QUEBRA
    base = len(AUDITORIA_PLANO) + 3
    cols = [("Data", 12), ("Tipo", 20), ("Realizada por", 18), ("Escopo / amostra", 40), ("Divergências", 40),
            ("Ação corretiva", 40), ("Concluída?", 10)]
    for i, (n, _) in enumerate(cols, 1):
        c = ws.cell(base, i, n)
        c.font, c.fill, c.alignment = BRANCO, preencher(AZUL), CENTRO
    for i, (_, w) in enumerate(cols, 1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = max(ws.column_dimensions[letra].width or 0, w)
    validacao(ws, [p for p, *_ in AUDITORIA_PLANO], f"B{base + 1}:B{base + 200}")
    validacao(ws, ["Sim", "Não"], f"G{base + 1}:G{base + 200}")
    for r in range(base + 1, base + 41):
        ws[f"A{r}"].number_format = DATA
    exec_base = base + 203
    ws.cell(exec_base, 1, "Registro de execuções do robô (preenchido automaticamente)").font = Font(bold=True, color=AZUL)
    for i, n in enumerate(["Data/hora", "Fonte (DJEN/DataJud)", "Publicações lidas", "Linhas novas", "Para revisão",
                           "Erros", "Versão"], 1):
        c = ws.cell(exec_base + 1, i, n)
        c.font, c.fill = BRANCO, preencher(CINZA)


def gerar(origem: Path, saida: Path, hoje: dt.date, cadastro: Path | None = None) -> dict:
    prazos, processos, decisoes, publicacoes, duplicatas = migrar(origem, hoje)
    clientes, avisos_cadastro = ler_cadastro(cadastro) if cadastro else ([], [])
    wb = Workbook()
    wb.remove(wb.active)
    aba_prazos(wb, prazos, hoje)
    combinacoes = sorted({(p["Tribunal"], p["Comarca"]) for p in prazos}, key=lambda x: (ordem_tribunal(x[0]), x[1]))
    aba_painel(wb, combinacoes, hoje)
    for trib in [t for t in ORDEM_TRIBUNAIS if any(p["Tribunal"] == t for p in prazos)]:
        aba_vista(wb, trib)
    aba_processos(wb, processos, clientes)
    aba_clientes(wb, clientes)
    aba_decisoes(wb, decisoes)
    aba_publicacoes(wb, publicacoes)
    aba_feriados(wb)
    aba_comarcas(wb)
    aba_auditoria(wb)
    aba_leia_me(wb)
    wb.save(saida)
    return {"prazos": len(prazos), "processos": len(processos), "decisoes": len(decisoes),
            "publicacoes": len(publicacoes), "duplicatas_removidas": duplicatas,
            "status": {s: sum(p["Status"] == s for p in prazos) for s in STATUS},
            "com_conferencia": sum(bool(p["Conferência"]) for p in prazos),
            "clientes": len(clientes), "avisos_cadastro": avisos_cadastro,
            "sugestoes_vinculo": sum(bool(p.get("Sugestão de vínculo (conferir)")) for p in processos)}


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Gera o modelo v2 da planilha e migra os dados da anterior.")
    p.add_argument("origem", type=Path)
    p.add_argument("-o", "--saida", type=Path, default=Path("modelo_v2.xlsx"))
    p.add_argument("--hoje", type=dt.date.fromisoformat, default=dt.date.today())
    p.add_argument("--cadastro", type=Path, help="exportação .xlsx do Cadastro Central de Clientes")
    a = p.parse_args(argv)
    print(gerar(a.origem, a.saida, a.hoje, a.cadastro))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
